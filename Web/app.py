import json
import os
import random
import re
import secrets
import sqlite3
import threading
import time
import unicodedata
import csv
import smtplib
import urllib.request
from email.message import EmailMessage
from zipfile import BadZipFile
from datetime import datetime, timezone, timedelta
from io import BytesIO, StringIO
from flask import Flask, g, jsonify, redirect, render_template, request, send_file, session
try:
    import serial
except ImportError:  # chạy không cần cổng Serial
    serial = None
import paho.mqtt.client as mqtt
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.middleware.proxy_fix import ProxyFix

load_dotenv(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env"))

APP_ENV = os.environ.get("APP_ENV", "development").strip().lower()
IS_PRODUCTION = APP_ENV in {"production", "prod"}
FLASK_SECRET_KEY = os.environ.get("FLASK_SECRET_KEY", "").strip()
if IS_PRODUCTION and len(FLASK_SECRET_KEY) < 32:
    raise RuntimeError("FLASK_SECRET_KEY phải được cấu hình bằng secret ngẫu nhiên dài ít nhất 32 ký tự khi APP_ENV=production.")
if not FLASK_SECRET_KEY:
    FLASK_SECRET_KEY = secrets.token_urlsafe(48)
    print("[Security] FLASK_SECRET_KEY chưa cấu hình; đã tạo secret tạm cho phiên chạy này.")

app = Flask(__name__)
app.secret_key = FLASK_SECRET_KEY
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SECURE=IS_PRODUCTION,
    SESSION_COOKIE_SAMESITE="Lax",
    PERMANENT_SESSION_LIFETIME=timedelta(hours=8),
    MAX_CONTENT_LENGTH=int(os.environ.get("MAX_CONTENT_LENGTH_MB", "8")) * 1024 * 1024,
    JSON_SORT_KEYS=False,
)
if IS_PRODUCTION:
    app.wsgi_app = ProxyFix(
        app.wsgi_app,
        x_for=int(os.environ.get("PROXY_FIX_X_FOR", "1")),
        x_proto=int(os.environ.get("PROXY_FIX_X_PROTO", "1")),
        x_host=int(os.environ.get("PROXY_FIX_X_HOST", "1")),
    )

# Cấu hình cổng Serial (Thay đổi COMx hoặc /dev/ttyUSB0 cho đúng với phần cứng của bạn)
SERIAL_PORT = os.environ.get("SERIAL_PORT", "COM3")
BAUD_RATE = int(os.environ.get("BAUD_RATE", "115200"))
DATABASE = os.environ.get("DATABASE_PATH", os.path.join(os.path.dirname(os.path.abspath(__file__)), "database.db"))
MQTT_BROKER = os.environ.get("MQTT_BROKER", "ha8969e1.ala.asia-southeast1.emqxsl.com")
MQTT_PORT = int(os.environ.get("MQTT_PORT", "8883"))
MQTT_TOPIC = os.environ.get("MQTT_TOPIC", "class/quiz/answer")
MQTT_USERNAME = os.environ.get("MQTT_USERNAME", "dang")
MQTT_PASSWORD = os.environ.get("MQTT_PASSWORD", "12345678")
MQTT_STATE = {
    "connected": False,
    "last_message_at": None,
    "last_message_result": None,
    "last_message_ok": None,
    "last_error": None,
    "messages_received": 0,
}
MQTT_TLS_INSECURE = os.environ.get("MQTT_TLS_INSECURE", "0") == "1"
if IS_PRODUCTION and MQTT_TLS_INSECURE:
    raise RuntimeError("MQTT_TLS_INSECURE=1 không được phép khi APP_ENV=production.")
MQTT_IS_V2 = hasattr(mqtt, "CallbackAPIVersion")  # paho-mqtt >= 2.0
MQTT_CONNECTED_EVENT = threading.Event()

# Khôi phục mật khẩu: gửi mã OTP qua email (SMTP) hoặc SMS (webhook của nhà cung cấp SMS).
SMTP_HOST = os.environ.get("SMTP_HOST", "")
SMTP_PORT = int(os.environ.get("SMTP_PORT", "587"))
SMTP_USER = os.environ.get("SMTP_USER", "")
SMTP_PASSWORD = os.environ.get("SMTP_PASSWORD", "")
SMTP_FROM = os.environ.get("SMTP_FROM", "") or SMTP_USER
SMS_WEBHOOK_URL = os.environ.get("SMS_WEBHOOK_URL", "")
SMS_WEBHOOK_TOKEN = os.environ.get("SMS_WEBHOOK_TOKEN", "")
# Nếu chưa cấu hình SMTP/SMS, in mã ra console của server (chỉ nên dùng khi chạy thử/nội bộ).
RESET_OTP_CONSOLE = os.environ.get("RESET_OTP_CONSOLE", "1") == "1"
if IS_PRODUCTION and RESET_OTP_CONSOLE:
    raise RuntimeError("RESET_OTP_CONSOLE=1 không được phép khi APP_ENV=production; OTP phải đi qua email/SMS.")
RESET_CODE_TTL_SECONDS = 10 * 60
RESET_MAX_ATTEMPTS = 5
RESET_MIN_INTERVAL_SECONDS = 60
RESET_MAX_PER_HOUR = 5

# -----------------------------------------------------------------------------
# Quản lý Cơ sở dữ liệu SQLite
# -----------------------------------------------------------------------------
def get_db_connection():
    conn = sqlite3.connect(DATABASE, timeout=20)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA busy_timeout = 20000")
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db():
    """Khởi tạo bảng cơ sở dữ liệu nếu chưa tồn tại"""
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("PRAGMA journal_mode = WAL")

    # Teacher-owned domain: accounts, classes, rooms and question content.
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS teacher_accounts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            email TEXT NOT NULL UNIQUE,
            password_hash TEXT NOT NULL,
            phone TEXT,
            created_at DATETIME DEFAULT CURRENT_TIMESTAMP
        )
    """
    )

    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS password_resets (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            teacher_id INTEGER NOT NULL,
            channel TEXT NOT NULL,
            code_hash TEXT NOT NULL,
            expires_at REAL NOT NULL,
            attempts INTEGER NOT NULL DEFAULT 0,
            used INTEGER NOT NULL DEFAULT 0,
            created_at REAL NOT NULL,
            FOREIGN KEY (teacher_id) REFERENCES teacher_accounts(id)
        )
    """
    )

    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS classes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            teacher_id INTEGER NOT NULL,
            name TEXT NOT NULL,
            description TEXT NOT NULL DEFAULT '',
            created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (teacher_id) REFERENCES teacher_accounts(id)
        )
    """
    )

    # Bảng danh sách sinh viên
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS students (
            student_id TEXT PRIMARY KEY,
            name TEXT NOT NULL,
            mac_address TEXT
        )
    """
    )

    # Phòng học và cấu hình mật khẩu truy cập
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS rooms (
            room_code TEXT PRIMARY KEY,
            password_hash TEXT NOT NULL,
            student_password_required INTEGER NOT NULL DEFAULT 1,
            active_question_id INTEGER,
            active_presentation_id INTEGER,
            question_published INTEGER NOT NULL DEFAULT 0,
            is_locked INTEGER NOT NULL DEFAULT 0,
            presentations_initialized INTEGER NOT NULL DEFAULT 0,
            created_at DATETIME DEFAULT CURRENT_TIMESTAMP
        )
    """
    )

    # Bảng nhật ký phản hồi / trả lời
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS responses (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            student_id TEXT NOT NULL,
            question_id INT NOT NULL,
            answer TEXT NOT NULL,
            timestamp DATETIME DEFAULT CURRENT_TIMESTAMP
        )
    """
    )

    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS questions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            room_code TEXT NOT NULL,
            position INTEGER NOT NULL,
            text TEXT NOT NULL,
            options_json TEXT NOT NULL,
            correct TEXT NOT NULL,
            explanation TEXT NOT NULL DEFAULT '',
            presentation_id INTEGER,
            source_question_id INTEGER,
            FOREIGN KEY (room_code) REFERENCES rooms(room_code),
            UNIQUE (room_code, position)
        )
    """
    )

    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS presentations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            room_code TEXT NOT NULL,
            title TEXT NOT NULL,
            scheduled_at TEXT,
            question_count INTEGER NOT NULL DEFAULT 0,
            status TEXT NOT NULL DEFAULT 'draft',
            is_default INTEGER NOT NULL DEFAULT 0,
            created_at REAL NOT NULL,
            opened_at REAL,
            closed_at REAL,
            block_id TEXT,
            FOREIGN KEY (room_code) REFERENCES rooms(room_code)
        )
        """
    )
    presentation_columns = {row["name"] for row in cursor.execute("PRAGMA table_info(presentations)")}
    if "block_id" not in presentation_columns:
        cursor.execute("ALTER TABLE presentations ADD COLUMN block_id TEXT")
    cursor.execute("CREATE INDEX IF NOT EXISTS idx_presentations_room ON presentations(room_code, id)")

    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS room_students (
            room_code TEXT NOT NULL,
            student_id TEXT NOT NULL,
            name TEXT NOT NULL,
            controller_id TEXT,
            roll_number INTEGER,
            joined_at DATETIME DEFAULT CURRENT_TIMESTAMP,
            last_seen DATETIME DEFAULT CURRENT_TIMESTAMP,
            PRIMARY KEY (room_code, student_id),
            FOREIGN KEY (room_code) REFERENCES rooms(room_code)
        )
    """
    )

    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS room_responses (
            room_code TEXT NOT NULL,
            question_id INTEGER NOT NULL,
            student_id TEXT NOT NULL,
            answer TEXT NOT NULL,
            timestamp DATETIME DEFAULT CURRENT_TIMESTAMP,
            PRIMARY KEY (room_code, question_id, student_id),
            FOREIGN KEY (room_code) REFERENCES rooms(room_code),
            FOREIGN KEY (question_id) REFERENCES questions(id)
        )
    """
    )

    cursor.execute(
        "CREATE INDEX IF NOT EXISTS idx_room_students_room ON room_students(room_code, joined_at)"
    )
    cursor.execute(
        "CREATE INDEX IF NOT EXISTS idx_room_responses_room_question ON room_responses(room_code, question_id, student_id)"
    )

    # Student-owned activity domain: room_students and room_responses below.
    # These stay in the same relational SQLite file so room/question references
    # and answer updates remain transactional, while SQL views separate reporting.
    # Read-only SQL views keep teacher-managed content and student activity distinct.
    cursor.execute(
        """
        CREATE VIEW IF NOT EXISTS teacher_class_data AS
        SELECT teachers.id AS teacher_id, teachers.name AS teacher_name,
               teachers.email, classes.id AS class_id, classes.name AS class_name,
               rooms.room_code, questions.id AS question_id,
               questions.position AS question_number, questions.text AS question_text
        FROM teacher_accounts AS teachers
        LEFT JOIN classes ON classes.teacher_id = teachers.id
        LEFT JOIN rooms ON rooms.class_id = classes.id AND rooms.teacher_id = teachers.id
        LEFT JOIN questions ON questions.room_code = rooms.room_code
        """
    )
    cursor.execute(
        """
        CREATE VIEW IF NOT EXISTS student_response_data AS
        SELECT room_students.room_code, room_students.student_id,
               room_students.name AS student_name, questions.id AS question_id,
               questions.position AS question_number, questions.text AS question_text,
               room_responses.answer, questions.correct AS correct_answer,
               CASE WHEN room_responses.answer IS NULL THEN 'unanswered'
                    WHEN room_responses.answer = questions.correct THEN 'correct'
                    ELSE 'incorrect' END AS result,
               room_responses.timestamp AS answered_at
        FROM room_students
        JOIN rooms ON rooms.room_code = room_students.room_code
        JOIN questions ON questions.room_code = rooms.room_code
        LEFT JOIN room_responses ON room_responses.room_code = room_students.room_code
             AND room_responses.student_id = room_students.student_id
             AND room_responses.question_id = questions.id
        """
    )

    # Kiểm tra trực tiếp trên app (độc lập với remote/máy chiếu):
    # mỗi bài kiểm tra có mã riêng; mỗi sinh viên có một lượt làm bài với bộ câu hỏi
    # và thứ tự đáp án được xáo trộn riêng, lưu cố định trong items_json.
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS exams (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            room_code TEXT NOT NULL,
            title TEXT NOT NULL,
            duration_minutes INTEGER NOT NULL,
            question_count INTEGER NOT NULL,
            access_code TEXT NOT NULL UNIQUE,
            status TEXT NOT NULL DEFAULT 'draft',
            show_answers INTEGER NOT NULL DEFAULT 0,
            created_at REAL NOT NULL,
            opened_at REAL,
            closed_at REAL,
            scheduled_at TEXT,
            question_pool_json TEXT,
            FOREIGN KEY (room_code) REFERENCES rooms(room_code)
        )
        """
    )
    exam_columns = {row["name"] for row in cursor.execute("PRAGMA table_info(exams)")}
    if "question_pool_json" not in exam_columns:
        cursor.execute("ALTER TABLE exams ADD COLUMN question_pool_json TEXT")
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS exam_attempts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            exam_id INTEGER NOT NULL,
            student_id TEXT NOT NULL,
            token TEXT NOT NULL UNIQUE,
            started_at REAL NOT NULL,
            deadline REAL NOT NULL,
            submitted_at REAL,
            submit_reason TEXT,
            items_json TEXT NOT NULL,
            correct_count INTEGER,
            total INTEGER,
            resume_allowed INTEGER NOT NULL DEFAULT 0,
            UNIQUE (exam_id, student_id),
            FOREIGN KEY (exam_id) REFERENCES exams(id)
        )
        """
    )
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS exam_answers (
            attempt_id INTEGER NOT NULL,
            question_id INTEGER NOT NULL,
            choice TEXT NOT NULL,
            answered_at REAL NOT NULL,
            PRIMARY KEY (attempt_id, question_id),
            FOREIGN KEY (attempt_id) REFERENCES exam_attempts(id)
        )
        """
    )
    cursor.execute("CREATE INDEX IF NOT EXISTS idx_exams_room ON exams(room_code, id)")

    # Mỗi tab trình duyệt có phiên riêng để tab học sinh không ghi đè tab giáo viên.
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS browser_sessions (
            session_id TEXT PRIMARY KEY,
            session_data TEXT NOT NULL,
            updated_at DATETIME DEFAULT CURRENT_TIMESTAMP
        )
    """
    )

    # Nâng cấp database được tạo bởi phiên bản cũ.
    room_columns = {row["name"] for row in cursor.execute("PRAGMA table_info(rooms)")}
    for column, definition in (
        ("active_question_id", "INTEGER"),
        ("active_presentation_id", "INTEGER"),
        ("question_published", "INTEGER NOT NULL DEFAULT 0"),
        ("is_locked", "INTEGER NOT NULL DEFAULT 0"),
        ("presentations_initialized", "INTEGER NOT NULL DEFAULT 0"),
        ("teacher_id", "INTEGER"),
        ("class_id", "INTEGER"),
        ("launched_at", "REAL"),
        ("block_id", "TEXT"),
    ):
        if column not in room_columns:
            cursor.execute(f"ALTER TABLE rooms ADD COLUMN {column} {definition}")

    question_columns = {row["name"] for row in cursor.execute("PRAGMA table_info(questions)")}
    for column in ("presentation_id", "source_question_id"):
        if column not in question_columns:
            cursor.execute(f"ALTER TABLE questions ADD COLUMN {column} INTEGER")

    # Migrate legacy room questions into a presentation while keeping the
    # original question bank intact for new presentations and app-based exams.
    for room_row in cursor.execute(
        "SELECT room_code, active_question_id, active_presentation_id, presentations_initialized, block_id "
        "FROM rooms"
    ).fetchall():
        room_code = room_row["room_code"]
        if not room_row["presentations_initialized"]:
            presentation = cursor.execute(
                "SELECT id FROM presentations WHERE room_code = ? AND is_default = 1 ORDER BY id LIMIT 1",
                (room_code,),
            ).fetchone()
            if presentation is None:
                source_questions = cursor.execute(
                    "SELECT id, position, text, options_json, correct, explanation FROM questions "
                    "WHERE room_code = ? AND presentation_id IS NULL ORDER BY position",
                    (room_code,),
                ).fetchall()
                if source_questions:
                    cursor.execute(
                        "INSERT INTO presentations (room_code, title, question_count, status, is_default, created_at) "
                        "VALUES (?, 'Trình chiếu hiện tại', 0, 'open', 1, ?)",
                        (room_code, time.time()),
                    )
                    presentation_id = cursor.lastrowid
                    next_position = cursor.execute(
                        "SELECT COALESCE(MAX(position), 0) FROM questions WHERE room_code = ?",
                        (room_code,),
                    ).fetchone()[0]
                    question_id_map = {}
                    for offset, question in enumerate(source_questions, start=1):
                        cloned = cursor.execute(
                            "INSERT INTO questions (room_code, position, text, options_json, correct, explanation, presentation_id, source_question_id) "
                            "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                            (
                                room_code,
                                next_position + offset,
                                question["text"],
                                question["options_json"],
                                question["correct"],
                                question["explanation"],
                                presentation_id,
                                question["id"],
                            ),
                        )
                        question_id_map[question["id"]] = cloned.lastrowid
                    for source_id, presentation_question_id in question_id_map.items():
                        cursor.execute(
                            "UPDATE room_responses SET question_id = ? WHERE room_code = ? AND question_id = ?",
                            (presentation_question_id, room_code, source_id),
                        )
                    if room_row["active_question_id"] in question_id_map:
                        cursor.execute(
                            "UPDATE rooms SET active_question_id = ? WHERE room_code = ?",
                            (question_id_map[room_row["active_question_id"]], room_code),
                        )
                    cursor.execute(
                        "UPDATE presentations SET question_count = ? WHERE id = ?",
                        (len(source_questions), presentation_id),
                    )
                    cursor.execute(
                        "UPDATE rooms SET active_presentation_id = ? WHERE room_code = ?",
                        (presentation_id, room_code),
                    )
            else:
                cursor.execute(
                    "UPDATE rooms SET active_presentation_id = ? "
                    "WHERE room_code = ? AND active_presentation_id IS NULL",
                    (presentation["id"], room_code),
                )
            cursor.execute(
                "UPDATE rooms SET presentations_initialized = 1 WHERE room_code = ?",
                (room_code,),
            )
        legacy_block_id = str(room_row["block_id"] or "").strip().upper()
        if legacy_block_id:
            active_presentation_id = cursor.execute(
                "SELECT active_presentation_id FROM rooms WHERE room_code = ?",
                (room_code,),
            ).fetchone()["active_presentation_id"]
            if active_presentation_id and re.fullmatch(r"[A-Z0-9_-]{1,32}", legacy_block_id):
                cursor.execute(
                    "UPDATE presentations SET block_id = ? WHERE id = ? AND block_id IS NULL",
                    (legacy_block_id, active_presentation_id),
                )
            cursor.execute("UPDATE rooms SET block_id = NULL WHERE room_code = ?", (room_code,))

    empty_default_presentations = cursor.execute(
        "SELECT id FROM presentations AS p WHERE p.is_default = 1 AND p.question_count = 0 "
        "AND NOT EXISTS (SELECT 1 FROM questions AS q WHERE q.presentation_id = p.id)"
    ).fetchall()
    if empty_default_presentations:
        empty_default_ids = [row["id"] for row in empty_default_presentations]
        placeholders = ",".join("?" for _ in empty_default_ids)
        cursor.execute(
            f"UPDATE rooms SET active_presentation_id = NULL, active_question_id = NULL, "
            f"question_published = 0, is_locked = 1, launched_at = NULL "
            f"WHERE active_presentation_id IN ({placeholders})",
            empty_default_ids,
        )
        cursor.execute(
            f"DELETE FROM presentations WHERE id IN ({placeholders})",
            empty_default_ids,
        )

    # Ngày kiểm tra của bài kiểm tra (database tạo từ phiên bản trước chưa có cột này).
    exam_columns = {row["name"] for row in cursor.execute("PRAGMA table_info(exams)")}
    if "scheduled_at" not in exam_columns:
        cursor.execute("ALTER TABLE exams ADD COLUMN scheduled_at TEXT")

    teacher_columns = {row["name"] for row in cursor.execute("PRAGMA table_info(teacher_accounts)")}
    if "phone" not in teacher_columns:
        cursor.execute("ALTER TABLE teacher_accounts ADD COLUMN phone TEXT")
    # Mỗi số điện thoại chỉ thuộc một tài khoản (tài khoản cũ chưa có số vẫn hợp lệ).
    cursor.execute(
        "CREATE UNIQUE INDEX IF NOT EXISTS idx_teacher_accounts_phone "
        "ON teacher_accounts(phone) WHERE phone IS NOT NULL"
    )

    roster_columns = {row["name"] for row in cursor.execute("PRAGMA table_info(room_students)")}
    for column, definition in (
        ("controller_id", "TEXT"),
        ("roll_number", "INTEGER"),
        ("note", "TEXT NOT NULL DEFAULT ''"),
    ):
        if column not in roster_columns:
            cursor.execute(f"ALTER TABLE room_students ADD COLUMN {column} {definition}")

    conn.commit()
    conn.close()


# Khởi tạo cả khi ứng dụng được chạy bằng WSGI (không qua khối __main__).
init_db()


@app.before_request
def load_tab_session():
    """Nạp phiên theo token riêng của tab, thay vì dùng chung cookie giữa các tab."""
    session_id = request.headers.get("X-Room-Session", "").strip()
    g.room_session_id = session_id or None
    if not session_id:
        return

    conn = get_db_connection()
    row = conn.execute(
        "SELECT session_data FROM browser_sessions WHERE session_id = ?", (session_id,)
    ).fetchone()
    conn.close()
    session.clear()
    if row is not None:
        try:
            stored_session = json.loads(row["session_data"])
            if isinstance(stored_session, dict):
                session.update(stored_session)
        except json.JSONDecodeError:
            pass


@app.after_request
def save_tab_session(response):
    """Lưu hoặc xóa phiên tab sau khi route hoàn tất."""
    session_id = getattr(g, "room_session_id", None)
    if not session_id:
        return response

    conn = get_db_connection()
    current_session = dict(session)
    if current_session.get("role") and (
        current_session.get("room_code") or current_session.get("teacher_id")
    ):
        conn.execute(
            "INSERT INTO browser_sessions (session_id, session_data, updated_at) "
            "VALUES (?, ?, CURRENT_TIMESTAMP) "
            "ON CONFLICT(session_id) DO UPDATE SET "
            "session_data = excluded.session_data, updated_at = CURRENT_TIMESTAMP",
            (session_id, json.dumps(current_session)),
        )
    else:
        conn.execute("DELETE FROM browser_sessions WHERE session_id = ?", (session_id,))
    conn.commit()
    conn.close()
    return response


# -----------------------------------------------------------------------------
# Security middleware
# -----------------------------------------------------------------------------
RATE_LIMIT_BUCKETS = {}
RATE_LIMIT_LOCK = threading.Lock()
RATE_LIMIT_RULES = {
    "/api/teacher/login": (10, 60),
    "/api/teacher/register": (5, 3600),
    "/api/teacher/password-reset/request": (5, 3600),
    "/api/teacher/password-reset/confirm": (10, 600),
    "/api/exam/lookup": (30, 60),
    "/api/exam/start": (20, 60),
}

def client_ip():
    return request.remote_addr or "unknown"

def rate_limited(key, limit, window_seconds):
    now = time.monotonic()
    bucket_key = (key, client_ip())
    with RATE_LIMIT_LOCK:
        timestamps = [t for t in RATE_LIMIT_BUCKETS.get(bucket_key, []) if now - t < window_seconds]
        if len(timestamps) >= limit:
            RATE_LIMIT_BUCKETS[bucket_key] = timestamps
            return True
        timestamps.append(now)
        RATE_LIMIT_BUCKETS[bucket_key] = timestamps
    return False

def same_origin_request():
    origin = request.headers.get("Origin", "").strip()
    if origin:
        return origin == request.host_url.rstrip("/")
    referer = request.headers.get("Referer", "").strip()
    if referer:
        return referer.startswith(request.host_url)
    return False

@app.before_request
def security_before_request():
    session.permanent = True
    if request.method in {"POST", "PUT", "PATCH", "DELETE"} and request.path.startswith("/api/"):
        if not same_origin_request():
            return jsonify({"status": "error", "message": "Yêu cầu không hợp lệ (CSRF/origin check)."}), 403
    rule = RATE_LIMIT_RULES.get(request.path)
    if rule and request.method == "POST":
        limit, window = rule
        if rate_limited(request.path, limit, window):
            response = jsonify({"status": "error", "message": "Quá nhiều yêu cầu. Vui lòng thử lại sau."})
            response.status_code = 429
            response.headers["Retry-After"] = str(window)
            return response

@app.after_request
def add_security_headers(response):
    response.headers.setdefault("X-Content-Type-Options", "nosniff")
    response.headers.setdefault("X-Frame-Options", "DENY")
    response.headers.setdefault("Referrer-Policy", "strict-origin-when-cross-origin")
    response.headers.setdefault("Permissions-Policy", "camera=(), microphone=(), geolocation=()")
    response.headers.setdefault("Cross-Origin-Opener-Policy", "same-origin")
    if request.path.startswith("/api/"):
        response.headers.setdefault("Cache-Control", "no-store")
    if IS_PRODUCTION and request.is_secure:
        response.headers.setdefault("Strict-Transport-Security", "max-age=31536000; includeSubDomains")
    return response

@app.errorhandler(413)
def request_too_large(_error):
    return jsonify({"status": "error", "message": "Tệp hoặc request vượt quá giới hạn cho phép."}), 413

# -----------------------------------------------------------------------------
# Luồng đọc dữ liệu Realtime từ Cổng Serial (ESP32 Gateway)
# -----------------------------------------------------------------------------
def save_response_to_db(student_id, question_id, answer):
    """Lưu phản hồi vào cơ sở dữ liệu"""
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute(
            "INSERT INTO responses (student_id, question_id, answer) VALUES (?, ?, ?)",
            (student_id, question_id, answer),
        )
        conn.commit()
        conn.close()
        print(f"[DB] Đã lưu: SV {student_id} - Câu {question_id} - ĐÁ: {answer}")
    except Exception as e:
        print(f"[DB Error] Lỗi lưu dữ liệu: {e}")


def read_serial_data():
    """Đọc gói tin JSON từ ESP32 qua Serial"""
    if serial is None:
        print("[Serial] Chưa cài pyserial; bỏ qua cổng Serial (vẫn nhận đáp án qua MQTT).")
        return
    try:
        ser = serial.Serial(SERIAL_PORT, BAUD_RATE, timeout=1)
        print(f"[Serial] Đã kết nối thành công với {SERIAL_PORT}")
        while True:
            if ser.in_waiting > 0:
                line = ser.readline().decode("utf-8", errors="ignore").strip()
                if line:
                    try:
                        data = json.loads(line)
                        student_id = data.get("student_id")
                        question_id = data.get("question_id")
                        answer = data.get("answer")

                        if student_id and question_id and answer:
                            save_response_to_db(student_id, question_id, answer)
                    except json.JSONDecodeError:
                        pass
            time.sleep(0.05)
    except Exception as e:
        print(f"[Serial Warning] Không thể mở cổng {SERIAL_PORT}: {e}")
        print("[Serial] Hệ thống vẫn chạy ở chế độ Web/Simulation.")


def normalize_mqtt_timestamp(value):
    """Convert Unix seconds/milliseconds or ISO timestamps to SQLite UTC time."""
    if value is None or value == "":
        return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
    try:
        timestamp = float(value)
        if timestamp > 10_000_000_000:
            timestamp /= 1000
        return datetime.fromtimestamp(timestamp, timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
    except (TypeError, ValueError, OverflowError, OSError):
        pass
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        return parsed.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
    except ValueError:
        return None


def normalize_block_id(value):
    """Chuẩn hóa ID khối remote (vd ' 333 ' -> '333'). Chỉ cho phép chữ, số, _ và -."""
    text = str(value or "").strip().upper()
    return text if re.fullmatch(r"[A-Z0-9_-]{1,32}", text) else ""


def ingest_mqtt_answer(payload, topic=MQTT_TOPIC):
    """Ghi đáp án từ remote vào phòng đang gán đúng ID khối.

    Định dạng khuyến nghị:  {"block_id": "333", "stt": 12, "answer": "B"}
    - block_id: ID khối remote được gán riêng cho bài trình chiếu đang hoạt động.
    - stt: số thứ tự (hoặc MSSV) của học sinh trong danh sách lớp.
    Nếu tin không có block_id thì dùng cách cũ: tự chọn phòng phát câu hỏi gần nhất.
    """
    if topic != MQTT_TOPIC:
        return False, "topic không khớp"
    if isinstance(payload, bytes):
        payload = payload.decode("utf-8")
    if isinstance(payload, str):
        payload = json.loads(payload)
    if not isinstance(payload, dict):
        return False, "JSON phải là object"

    block_id = normalize_block_id(payload.get("block_id", payload.get("block", payload.get("kit_id"))))
    controller_id = str(
        payload.get(
            "stt",
            payload.get("roll_number", payload.get("controller_id", payload.get("device_id", payload.get("student_id", "")))),
        )
    ).strip()
    answer = str(payload.get("answer", "")).strip().upper()
    room_code = str(payload.get("room_code", "")).strip().upper()
    requested_question_id = payload.get("question_id")
    timestamp = normalize_mqtt_timestamp(payload.get("timestamp"))
    if not controller_id or answer not in {"A", "B", "C", "D"}:
        return False, "cần stt (hoặc student_id) hợp lệ và answer A/B/C/D"
    if timestamp is None:
        return False, "timestamp không hợp lệ"

    conn = get_db_connection()
    query = (
        "SELECT rs.room_code, rs.student_id, r.active_question_id, r.launched_at, "
        "(CAST(rs.roll_number AS TEXT) = ?) AS by_roll "
        "FROM room_students AS rs JOIN rooms AS r ON r.room_code = rs.room_code "
        "JOIN presentations AS p ON p.id = r.active_presentation_id AND p.room_code = r.room_code "
        "WHERE (CAST(rs.roll_number AS TEXT) = ? OR rs.student_id = ?) "
        "AND r.question_published = 1 AND r.is_locked = 0 "
        "AND r.active_question_id IS NOT NULL"
    )
    parameters = [controller_id, controller_id, controller_id]
    if block_id:
        query += " AND UPPER(p.block_id) = ?"
        parameters.append(block_id)
    if room_code:
        query += " AND rs.room_code = ?"
        parameters.append(room_code)
    if requested_question_id not in (None, ""):
        try:
            requested_question_id = int(requested_question_id)
        except (TypeError, ValueError):
            conn.close()
            return False, "question_id không hợp lệ"
        query += " AND r.active_question_id = ?"
        parameters.append(requested_question_id)

    matches = conn.execute(query, parameters).fetchall()
    if not matches:
        # Chẩn đoán lý do để giảng viên biết cần sửa gì.
        if block_id:
            room_row = conn.execute(
                "SELECT r.room_code, r.question_published, r.is_locked, r.active_question_id "
                "FROM presentations AS p JOIN rooms AS r "
                "ON r.room_code = p.room_code AND r.active_presentation_id = p.id "
                "WHERE UPPER(p.block_id) = ?",
                (block_id,),
            ).fetchone()
            if room_row is None:
                conn.close()
                return False, f"khối '{block_id}' chưa được gán cho bài trình chiếu đang mở"
            rooms_known = [room_row]
        else:
            diag_sql = (
                "SELECT rs.room_code, r.question_published, r.is_locked, r.active_question_id "
                "FROM room_students AS rs JOIN rooms AS r ON r.room_code = rs.room_code "
                "WHERE (CAST(rs.roll_number AS TEXT) = ? OR rs.student_id = ?)"
            )
            diag_params = [controller_id, controller_id]
            if room_code:
                diag_sql += " AND rs.room_code = ?"
                diag_params.append(room_code)
            rooms_known = conn.execute(diag_sql, diag_params).fetchall()
            if not rooms_known:
                conn.close()
                where = f" của phòng {room_code}" if room_code else ""
                return False, f"ID '{controller_id}' không có trong danh sách lớp{where}"
        row = rooms_known[0]
        if not row["question_published"] or row["active_question_id"] is None:
            conn.close()
            return False, f"phòng {row['room_code']} chưa phát câu hỏi"
        if row["is_locked"]:
            conn.close()
            return False, f"phòng {row['room_code']} đang bị khóa"
        if block_id:
            in_roster = conn.execute(
                "SELECT 1 FROM room_students WHERE room_code = ? "
                "AND (CAST(roll_number AS TEXT) = ? OR student_id = ?)",
                (row["room_code"], controller_id, controller_id),
            ).fetchone()
            conn.close()
            if in_roster is None:
                return False, f"STT '{controller_id}' không có trong danh sách lớp của phòng {row['room_code']}"
            return False, "question_id trong tin nhắn không khớp câu hỏi đang phát"
        conn.close()
        return False, "question_id trong tin nhắn không khớp câu hỏi đang phát"
    if len(matches) > 1:
        # Không có block_id: ưu tiên phòng vừa phát gần nhất, rồi ưu tiên khớp theo STT.
        newest = max((row["launched_at"] or 0) for row in matches)
        matches = [row for row in matches if (row["launched_at"] or 0) == newest]
        if len(matches) > 1 and any(row["by_roll"] for row in matches):
            matches = [row for row in matches if row["by_roll"]]
        if len({(row["room_code"], row["student_id"]) for row in matches}) > 1:
            conn.close()
            rooms_hit = ", ".join(sorted({str(row["room_code"]) for row in matches}))
            return False, (
                f"ID '{controller_id}' vẫn khớp nhiều học sinh/phòng ({rooms_hit}). "
                "Hãy gán Khối ID cho phòng hoặc kiểm tra STT/MSSV bị trùng trong danh sách lớp"
            )

    match = matches[0]
    student_id = match["student_id"]
    question = conn.execute(
        "SELECT id FROM questions WHERE id = ? AND room_code = ?",
        (match["active_question_id"], match["room_code"]),
    ).fetchone()
    if question is None:
        conn.close()
        return False, "không tìm thấy câu hỏi đang phát trong phòng"

    conn.execute(
        "INSERT INTO room_responses (room_code, question_id, student_id, answer, timestamp) "
        "VALUES (?, ?, ?, ?, ?) "
        "ON CONFLICT(room_code, question_id, student_id) DO UPDATE SET "
        "answer = excluded.answer, timestamp = excluded.timestamp",
        (match["room_code"], question["id"], student_id, answer, timestamp),
    )
    conn.execute(
        "UPDATE room_students SET last_seen = CURRENT_TIMESTAMP "
        "WHERE room_code = ? AND student_id = ?",
        (match["room_code"], student_id),
    )
    conn.commit()
    conn.close()
    block_note = f" (khối {block_id})" if block_id else ""
    return True, f"đã ghi nhận {student_id} / {answer} cho phòng {match['room_code']}{block_note}"


def create_mqtt_client(client_id):
    """Tạo client tương thích cả paho-mqtt 1.x và 2.x."""
    if MQTT_IS_V2:
        return mqtt.Client(
            callback_api_version=mqtt.CallbackAPIVersion.VERSION2,
            client_id=client_id,
            protocol=mqtt.MQTTv311,
            transport="tcp",
        )
    return mqtt.Client(client_id=client_id, protocol=mqtt.MQTTv311, transport="tcp")


def start_mqtt_client():
    """Kết nối EMQX qua MQTT over TLS (cổng 8883), tự kết nối lại khi mất mạng."""
    if not MQTT_USERNAME or not MQTT_PASSWORD:
        MQTT_STATE["last_error"] = "Thiếu MQTT_USERNAME hoặc MQTT_PASSWORD trong file .env."
        print("[MQTT] Chưa cấu hình MQTT_USERNAME/MQTT_PASSWORD; bỏ qua kết nối broker.")
        return None

    client = create_mqtt_client(f"ars-server-{secrets.token_hex(4)}")
    client.username_pw_set(MQTT_USERNAME, MQTT_PASSWORD)
    if MQTT_TLS_INSECURE:
        import ssl

        context = ssl.create_default_context()
        context.check_hostname = False
        context.verify_mode = ssl.CERT_NONE
        client.tls_set_context(context)
    else:
        client.tls_set()
    client.reconnect_delay_set(min_delay=1, max_delay=15)

    def on_connect(mqtt_client, _userdata, _flags, reason_code, _properties=None):
        if reason_code != 0:
            MQTT_STATE["connected"] = False
            MQTT_CONNECTED_EVENT.clear()
            MQTT_STATE["last_error"] = f"Broker từ chối kết nối: {reason_code} (kiểm tra user/password trên EMQX)"
            print(f"[MQTT] Broker từ chối kết nối: {reason_code}")
            return
        # Subscribe trong on_connect để tự đăng ký lại sau mỗi lần kết nối lại.
        result, _message_id = mqtt_client.subscribe(MQTT_TOPIC, qos=1)
        if result != mqtt.MQTT_ERR_SUCCESS:
            MQTT_STATE["connected"] = False
            MQTT_STATE["last_error"] = f"Không subscribe được topic {MQTT_TOPIC}: {result}"
            print(f"[MQTT] Subscribe thất bại: {result}")
            return
        MQTT_STATE["connected"] = True
        MQTT_CONNECTED_EVENT.set()
        MQTT_STATE["last_error"] = None
        print(f"[MQTT] Đã kết nối {MQTT_BROKER}:{MQTT_PORT}, subscribe {MQTT_TOPIC}")

    def on_disconnect(_mqtt_client, _userdata, *args):
        # v2: (disconnect_flags, reason_code, properties) | v1: (rc,)
        reason_code = args[1] if MQTT_IS_V2 and len(args) > 1 else args[0]
        MQTT_STATE["connected"] = False
        MQTT_CONNECTED_EVENT.clear()
        if reason_code != 0:
            MQTT_STATE["last_error"] = f"MQTT bị ngắt: {reason_code} (đang tự kết nối lại)"
            print(f"[MQTT] Kết nối bị ngắt: {reason_code}; đang tự kết nối lại...")

    def on_connect_fail(_mqtt_client, _userdata):
        MQTT_STATE["connected"] = False
        MQTT_STATE["last_error"] = (
            f"Không kết nối được {MQTT_BROKER}:{MQTT_PORT} - kiểm tra mạng/tường lửa (cổng 8883)"
        )
        print(f"[MQTT Warning] {MQTT_STATE['last_error']}")

    def on_message(_mqtt_client, _userdata, message):
        MQTT_STATE["messages_received"] += 1
        MQTT_STATE["last_message_at"] = datetime.now(timezone.utc).isoformat(timespec="seconds")
        print(f"[MQTT] Nhận bản tin trên topic {message.topic}")
        try:
            success, _detail = ingest_mqtt_answer(message.payload, message.topic)
            MQTT_STATE["last_message_result"] = (
                "Đã ghi nhận đáp án." if success else "Bản tin không được ghi nhận."
            )
            MQTT_STATE["last_message_ok"] = success
            MQTT_STATE["last_error"] = None if success else "Bản tin MQTT bị từ chối."
            print("[MQTT] Đã xử lý bản tin." if success else "[MQTT Warning] Bản tin bị từ chối.")
        except json.JSONDecodeError as error:
            MQTT_STATE["last_message_ok"] = False
            MQTT_STATE["last_message_result"] = f"Payload JSON không hợp lệ: {error}"
            MQTT_STATE["last_error"] = MQTT_STATE["last_message_result"]
            print(f"[MQTT Warning] {MQTT_STATE['last_message_result']}")
        except Exception:
            MQTT_STATE["last_message_ok"] = False
            MQTT_STATE["last_message_result"] = "Không thể xử lý bản tin MQTT."
            MQTT_STATE["last_error"] = MQTT_STATE["last_message_result"]
            print(f"[MQTT Error] {MQTT_STATE['last_message_result']}")

    client.on_connect = on_connect
    client.on_disconnect = on_disconnect
    client.on_connect_fail = on_connect_fail
    client.on_message = on_message
    try:
        # connect_async: không bị dừng hẳn khi lúc khởi động chưa có mạng, tự thử lại.
        client.connect_async(MQTT_BROKER, MQTT_PORT, keepalive=60)
        client.loop_start()
    except Exception as error:
        MQTT_STATE["connected"] = False
        MQTT_STATE["last_error"] = f"Không khởi động được MQTT: {error}"
        print(f"[MQTT Warning] {MQTT_STATE['last_error']}")
        return None
    return client


# -----------------------------------------------------------------------------
# Routes API & Web Interface
# -----------------------------------------------------------------------------
@app.route("/")
def index():
    """Mở trang giao diện Dashboard"""
    return render_template("index.html", page_role=None)


@app.route("/health/live")
@app.route("/healthz")
def health_live():
    return jsonify({"status": "alive"})


@app.route("/health/ready")
@app.route("/readyz")
def health_ready():
    conn = None
    try:
        conn = get_db_connection()
        conn.execute("SELECT 1").fetchone()
    except sqlite3.Error:
        app.logger.exception("Database readiness check failed")
        return jsonify({"status": "not_ready", "checks": {"database": "unavailable"}}), 503
    finally:
        if conn is not None:
            conn.close()
    return jsonify({"status": "ready", "checks": {"database": "ok"}})


@app.route("/privacy")
def privacy_policy():
    return render_template("privacy.html")


@app.route("/teacher")
def teacher_page():
    """Trang đăng nhập và dashboard dành riêng cho giảng viên."""
    return render_template("index.html", page_role="teacher")


@app.route("/student")
def student_page():
    """Chuyển luồng học sinh cũ về trang giảng viên; clicker không đăng nhập web."""
    return redirect("/teacher")


@app.route("/exam")
@app.route("/exam/<access_code>")
def exam_page(access_code=""):
    """Trang làm bài kiểm tra dành cho sinh viên (không cần đăng nhập tài khoản)."""
    return render_template("exam.html", prefill_code=access_code.strip().upper())


@app.route("/projector/<room_code>")
def projector(room_code):
    """Màn hình trình chiếu công khai, cập nhật theo câu hỏi được phát."""
    return render_template("projector.html", room_code=room_code.strip().upper())


def get_json_data():
    """Đọc JSON an toàn; trả về dict rỗng nếu request không hợp lệ."""
    data = request.get_json(silent=True)
    return data if isinstance(data, dict) else {}


def teacher_account_required():
    return session.get("role") == "teacher" and bool(session.get("teacher_id"))


def normalize_phone(value):
    """Chuẩn hóa số điện thoại di động Việt Nam về dạng 0xxxxxxxxx; trả về '' nếu không hợp lệ."""
    digits = re.sub(r"[\s.\-()]", "", str(value or ""))
    if digits.startswith("+84"):
        digits = "0" + digits[3:]
    elif digits.startswith("84") and len(digits) == 11:
        digits = "0" + digits[2:]
    return digits if re.fullmatch(r"0(?:3|5|7|8|9)\d{8}", digits) else ""


def deliver_reset_code(channel, destination, code):
    """Gửi mã OTP qua email/SMS. Chạy ở luồng riêng để thời gian phản hồi không lộ tài khoản có tồn tại hay không."""
    message = (
        f"Ma dat lai mat khau HCMUTE ARS cua ban la {code}. "
        "Ma co hieu luc 10 phut. Khong chia se ma nay cho bat ky ai."
    )
    try:
        if channel == "email" and SMTP_HOST:
            mail = EmailMessage()
            mail["Subject"] = "Mã đặt lại mật khẩu HCMUTE ARS"
            mail["From"] = SMTP_FROM
            mail["To"] = destination
            mail.set_content(
                f"Mã đặt lại mật khẩu của bạn là: {code}\n\n"
                "Mã có hiệu lực trong 10 phút. Nếu bạn không yêu cầu, hãy bỏ qua email này "
                "và không chia sẻ mã cho bất kỳ ai."
            )
            with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=15) as server:
                server.starttls()
                if SMTP_USER:
                    server.login(SMTP_USER, SMTP_PASSWORD)
                server.send_message(mail)
            return
        if channel == "sms" and SMS_WEBHOOK_URL:
            international = "+84" + destination[1:]
            request_headers = {"Content-Type": "application/json"}
            if SMS_WEBHOOK_TOKEN:
                request_headers["Authorization"] = f"Bearer {SMS_WEBHOOK_TOKEN}"
            sms_request = urllib.request.Request(
                SMS_WEBHOOK_URL,
                data=json.dumps({"phone": international, "message": message}).encode("utf-8"),
                headers=request_headers,
                method="POST",
            )
            with urllib.request.urlopen(sms_request, timeout=15):
                pass
            return
    except Exception as error:
        print(f"[Reset] Gửi mã qua {channel} thất bại: {error}")
    if RESET_OTP_CONSOLE:
        print(f"[Reset OTP] Kênh {channel} -> {destination}: mã {code} (hết hạn sau 10 phút)")


def find_teacher_by_identifier(conn, identifier):
    """Tìm giảng viên theo email hoặc số điện thoại. Trả về (channel, teacher_row | None); channel rỗng nếu định dạng sai."""
    identifier = str(identifier or "").strip()
    if "@" in identifier:
        row = conn.execute(
            "SELECT id, name, email, phone FROM teacher_accounts WHERE email = ?",
            (identifier.lower(),),
        ).fetchone()
        return "email", row
    phone = normalize_phone(identifier)
    if not phone:
        return "", None
    row = conn.execute(
        "SELECT id, name, email, phone FROM teacher_accounts WHERE phone = ?", (phone,)
    ).fetchone()
    return "sms", row


def revoke_teacher_sessions(conn, teacher_id):
    """Đăng xuất mọi tab/thiết bị của giảng viên (dùng sau khi đổi mật khẩu)."""
    for browser_session in conn.execute("SELECT session_id, session_data FROM browser_sessions").fetchall():
        try:
            data = json.loads(browser_session["session_data"])
        except json.JSONDecodeError:
            continue
        if data.get("teacher_id") == teacher_id:
            conn.execute("DELETE FROM browser_sessions WHERE session_id = ?", (browser_session["session_id"],))


@app.route("/api/teacher/register", methods=["POST"])
def register_teacher():
    data = get_json_data()
    name = str(data.get("name", "")).strip()
    email = str(data.get("email", "")).strip().lower()
    password = str(data.get("password", ""))
    phone = normalize_phone(data.get("phone"))
    if not name or "@" not in email or not password:
        return jsonify({"status": "error", "message": "Vui lòng nhập họ tên, email hợp lệ và mật khẩu."}), 400
    if not phone:
        return jsonify({"status": "error", "message": "Số điện thoại không hợp lệ (ví dụ: 0912345678)."}), 400
    if len(password) < 8:
        return jsonify({"status": "error", "message": "Mật khẩu tài khoản phải có ít nhất 8 ký tự."}), 400

    conn = get_db_connection()
    if conn.execute("SELECT 1 FROM teacher_accounts WHERE phone = ?", (phone,)).fetchone():
        conn.close()
        return jsonify({"status": "error", "message": "Số điện thoại này đã được đăng ký cho tài khoản khác."}), 409
    try:
        cursor = conn.execute(
            "INSERT INTO teacher_accounts (name, email, password_hash, phone) VALUES (?, ?, ?, ?)",
            (name, email, generate_password_hash(password), phone),
        )
        conn.commit()
        teacher_id = cursor.lastrowid
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({"status": "error", "message": "Email này đã được đăng ký."}), 409
    conn.close()

    session.clear()
    session.update({"role": "teacher", "teacher_id": teacher_id, "teacher_name": name})
    return jsonify({"status": "success", "teacher": {"id": teacher_id, "name": name, "email": email, "phone": phone}}), 201


@app.route("/api/teacher/password-reset/request", methods=["POST"])
def request_password_reset():
    """Bước 1: nhập email hoặc số điện thoại -> gửi mã OTP 6 số."""
    identifier = str(get_json_data().get("identifier", "")).strip()
    channel_hint = "email" if "@" in identifier else "sms"
    if not identifier or (channel_hint == "sms" and not normalize_phone(identifier)):
        return jsonify({"status": "error", "message": "Nhập email hoặc số điện thoại đã đăng ký."}), 400

    # Luôn trả cùng một thông báo để không lộ email/số điện thoại nào đã có tài khoản.
    generic = jsonify(
        {
            "status": "success",
            "channel": channel_hint,
            "message": "Nếu thông tin khớp với một tài khoản, mã xác nhận gồm 6 số đã được gửi. Mã có hiệu lực 10 phút.",
        }
    )
    conn = get_db_connection()
    channel, teacher = find_teacher_by_identifier(conn, identifier)
    if teacher is None:
        conn.close()
        return generic

    now = time.time()
    recent = conn.execute(
        "SELECT COUNT(*) AS total, MAX(created_at) AS last_at FROM password_resets "
        "WHERE teacher_id = ? AND created_at > ?",
        (teacher["id"], now - 3600),
    ).fetchone()
    if recent["total"] >= RESET_MAX_PER_HOUR or (
        recent["last_at"] and now - recent["last_at"] < RESET_MIN_INTERVAL_SECONDS
    ):
        conn.close()
        return generic

    code = f"{secrets.randbelow(1_000_000):06d}"
    conn.execute("UPDATE password_resets SET used = 1 WHERE teacher_id = ? AND used = 0", (teacher["id"],))
    conn.execute(
        "INSERT INTO password_resets (teacher_id, channel, code_hash, expires_at, created_at) VALUES (?, ?, ?, ?, ?)",
        (teacher["id"], channel, generate_password_hash(code), now + RESET_CODE_TTL_SECONDS, now),
    )
    conn.commit()
    conn.close()

    destination = teacher["email"] if channel == "email" else teacher["phone"]
    threading.Thread(target=deliver_reset_code, args=(channel, destination, code), daemon=True).start()
    return generic


@app.route("/api/teacher/password-reset/confirm", methods=["POST"])
def confirm_password_reset():
    """Bước 2: nhập mã OTP + mật khẩu mới."""
    data = get_json_data()
    identifier = str(data.get("identifier", "")).strip()
    code = str(data.get("code", "")).strip()
    new_password = str(data.get("new_password", ""))
    bad_code = jsonify({"status": "error", "message": "Mã không đúng hoặc đã hết hạn. Hãy yêu cầu mã mới."})
    if len(new_password) < 8:
        return jsonify({"status": "error", "message": "Mật khẩu mới phải có ít nhất 8 ký tự."}), 400
    if not re.fullmatch(r"\d{6}", code):
        return jsonify({"status": "error", "message": "Mã xác nhận gồm 6 chữ số."}), 400

    conn = get_db_connection()
    channel, teacher = find_teacher_by_identifier(conn, identifier)
    if teacher is None:
        conn.close()
        return bad_code, 400
    reset = conn.execute(
        "SELECT id, channel, code_hash, expires_at, attempts FROM password_resets "
        "WHERE teacher_id = ? AND used = 0 ORDER BY id DESC LIMIT 1",
        (teacher["id"],),
    ).fetchone()
    if (
        reset is None
        or reset["channel"] != channel
        or reset["expires_at"] < time.time()
        or reset["attempts"] >= RESET_MAX_ATTEMPTS
    ):
        conn.close()
        return bad_code, 400

    if not check_password_hash(reset["code_hash"], code):
        attempts = reset["attempts"] + 1
        conn.execute(
            "UPDATE password_resets SET attempts = ?, used = ? WHERE id = ?",
            (attempts, int(attempts >= RESET_MAX_ATTEMPTS), reset["id"]),
        )
        conn.commit()
        conn.close()
        return bad_code, 400

    conn.execute("BEGIN IMMEDIATE")
    try:
        conn.execute(
            "UPDATE teacher_accounts SET password_hash = ? WHERE id = ?",
            (generate_password_hash(new_password), teacher["id"]),
        )
        conn.execute("UPDATE password_resets SET used = 1 WHERE teacher_id = ?", (teacher["id"],))
        revoke_teacher_sessions(conn, teacher["id"])
        conn.commit()
    except Exception:
        conn.rollback()
        conn.close()
        raise
    conn.close()

    result = {"status": "success", "message": "Đã đặt lại mật khẩu. Hãy đăng nhập bằng mật khẩu mới."}
    if channel == "sms":
        # Đã xác minh quyền sở hữu số điện thoại nên có thể nhắc lại email đăng nhập cho người quên email.
        result["email"] = teacher["email"]
    return jsonify(result)


@app.route("/api/teacher/login", methods=["POST"])
def login_teacher():
    data = get_json_data()
    email = str(data.get("email", "")).strip().lower()
    password = str(data.get("password", ""))
    if not email or not password:
        return jsonify({"status": "error", "message": "Vui lòng nhập email và mật khẩu."}), 400

    conn = get_db_connection()
    teacher = conn.execute(
        "SELECT id, name, email, password_hash FROM teacher_accounts WHERE email = ?",
        (email,),
    ).fetchone()
    conn.close()
    if teacher is None or not check_password_hash(teacher["password_hash"], password):
        return jsonify({"status": "error", "message": "Email hoặc mật khẩu không đúng."}), 401

    session.clear()
    session.update({"role": "teacher", "teacher_id": teacher["id"], "teacher_name": teacher["name"]})
    return jsonify({"status": "success", "teacher": {"id": teacher["id"], "name": teacher["name"], "email": teacher["email"]}})


@app.route("/api/teacher/me", methods=["GET"])
def get_teacher_profile():
    if not teacher_account_required():
        return jsonify({"status": "error", "message": "Chưa đăng nhập tài khoản giảng viên."}), 401
    conn = get_db_connection()
    teacher = conn.execute(
        "SELECT id, name, email, phone FROM teacher_accounts WHERE id = ?",
        (session["teacher_id"],),
    ).fetchone()
    conn.close()
    if teacher is None:
        session.clear()
        return jsonify({"status": "error", "message": "Tài khoản không còn tồn tại."}), 401
    return jsonify({"status": "success", "teacher": dict(teacher)})


@app.route("/api/teacher/mqtt-status", methods=["GET"])
def get_teacher_mqtt_status():
    if not teacher_account_required():
        return jsonify({"status": "error", "message": "Vui lòng đăng nhập tài khoản giảng viên."}), 401
    result = {
        "connected": MQTT_STATE["connected"],
        "last_message_at": MQTT_STATE["last_message_at"],
        "last_message_ok": MQTT_STATE["last_message_ok"],
        "messages_received": MQTT_STATE["messages_received"],
        "last_message_result": (
            "Đã nhận và xử lý bản tin."
            if MQTT_STATE["last_message_ok"] is True
            else "Bản tin gần nhất không được xử lý."
            if MQTT_STATE["last_message_ok"] is False
            else None
        ),
    }
    return jsonify(
        {
            "status": "success",
            "broker": MQTT_BROKER,
            "port": MQTT_PORT,
            "topic": MQTT_TOPIC,
            **result,
        }
    )


@app.route("/api/teacher/logout", methods=["POST"])
def logout_teacher():
    session.clear()
    return jsonify({"status": "success"})


@app.route("/api/teacher/classes", methods=["GET", "POST"])
def teacher_classes():
    if not teacher_account_required():
        return jsonify({"status": "error", "message": "Vui lòng đăng nhập tài khoản giảng viên."}), 401

    conn = get_db_connection()
    if request.method == "POST":
        data = get_json_data()
        name = str(data.get("name", "")).strip()
        description = str(data.get("description", "")).strip()
        if not name:
            conn.close()
            return jsonify({"status": "error", "message": "Vui lòng nhập tên lớp."}), 400
        cursor = conn.execute(
            "INSERT INTO classes (teacher_id, name, description) VALUES (?, ?, ?)",
            (session["teacher_id"], name, description),
        )
        conn.commit()
        class_id = cursor.lastrowid
        conn.close()
        return jsonify({"status": "success", "class": {"id": class_id, "name": name, "description": description, "rooms": []}}), 201

    classes = conn.execute(
        "SELECT id, name, description, created_at FROM classes "
        "WHERE teacher_id = ? ORDER BY created_at DESC, id DESC",
        (session["teacher_id"],),
    ).fetchall()
    rooms = conn.execute(
        "SELECT class_id, room_code, student_password_required, created_at "
        "FROM rooms WHERE teacher_id = ? AND class_id IS NOT NULL "
        "ORDER BY created_at DESC",
        (session["teacher_id"],),
    ).fetchall()
    conn.close()
    rooms_by_class = {}
    for room in rooms:
        rooms_by_class.setdefault(room["class_id"], []).append(
            {
                "room_code": room["room_code"],
                "student_password_required": bool(room["student_password_required"]),
                "created_at": room["created_at"],
            }
        )
    return jsonify(
        {
            "status": "success",
            "classes": [
                {**dict(class_row), "rooms": rooms_by_class.get(class_row["id"], [])}
                for class_row in classes
            ],
        }
    )


def revoke_room_tab_sessions(conn, room_codes):
    """Remove deleted-room state from stored browser tabs, preserving teacher login."""
    room_codes = set(room_codes)
    if not room_codes:
        return
    sessions = conn.execute("SELECT session_id, session_data FROM browser_sessions").fetchall()
    for browser_session in sessions:
        try:
            data = json.loads(browser_session["session_data"])
        except json.JSONDecodeError:
            continue
        if data.get("room_code") not in room_codes:
            continue
        if data.get("role") == "teacher" and data.get("teacher_id"):
            data.pop("room_code", None)
            conn.execute(
                "UPDATE browser_sessions SET session_data = ?, updated_at = CURRENT_TIMESTAMP WHERE session_id = ?",
                (json.dumps(data), browser_session["session_id"]),
            )
        else:
            conn.execute(
                "DELETE FROM browser_sessions WHERE session_id = ?",
                (browser_session["session_id"],),
            )


@app.route("/api/teacher/classes/<int:class_id>", methods=["DELETE"])
def delete_teacher_class(class_id):
    if not teacher_account_required():
        return jsonify({"status": "error", "message": "Vui lòng đăng nhập tài khoản giảng viên."}), 401

    conn = get_db_connection()
    class_row = conn.execute(
        "SELECT id FROM classes WHERE id = ? AND teacher_id = ?",
        (class_id, session["teacher_id"]),
    ).fetchone()
    if class_row is None:
        conn.close()
        return jsonify({"status": "error", "message": "Không tìm thấy lớp của tài khoản này."}), 404

    room_rows = conn.execute(
        "SELECT room_code FROM rooms WHERE class_id = ? AND teacher_id = ?",
        (class_id, session["teacher_id"]),
    ).fetchall()
    room_codes = [row["room_code"] for row in room_rows]
    conn.execute("BEGIN IMMEDIATE")
    try:
        if room_codes:
            placeholders = ",".join("?" for _ in room_codes)
            delete_exam_data_for_rooms(conn, room_codes)
            conn.execute(f"DELETE FROM room_responses WHERE room_code IN ({placeholders})", room_codes)
            conn.execute(f"DELETE FROM room_students WHERE room_code IN ({placeholders})", room_codes)
            conn.execute(f"DELETE FROM presentations WHERE room_code IN ({placeholders})", room_codes)
            conn.execute(f"DELETE FROM questions WHERE room_code IN ({placeholders})", room_codes)
            conn.execute(
                "DELETE FROM rooms WHERE teacher_id = ? AND class_id = ?",
                (session["teacher_id"], class_id),
            )
        conn.execute("DELETE FROM classes WHERE id = ? AND teacher_id = ?", (class_id, session["teacher_id"]))
        revoke_room_tab_sessions(conn, room_codes)
        conn.commit()
    except Exception:
        conn.rollback()
        conn.close()
        raise
    conn.close()
    if session.get("room_code") in room_codes:
        session.pop("room_code", None)
    return jsonify({"status": "success", "message": "Đã xóa lớp và toàn bộ phòng, danh sách, câu hỏi, đáp án thuộc lớp."})


@app.route("/api/teacher/rooms/<room_code>", methods=["DELETE"])
def delete_teacher_room(room_code):
    if not teacher_account_required():
        return jsonify({"status": "error", "message": "Vui lòng đăng nhập tài khoản giảng viên."}), 401
    room_code = room_code.strip().upper()
    conn = get_db_connection()
    room = conn.execute(
        "SELECT room_code FROM rooms WHERE room_code = ? AND teacher_id = ?",
        (room_code, session["teacher_id"]),
    ).fetchone()
    if room is None:
        conn.close()
        return jsonify({"status": "error", "message": "Không tìm thấy phòng của tài khoản này."}), 404

    conn.execute("BEGIN IMMEDIATE")
    try:
        delete_exam_data_for_rooms(conn, [room_code])
        conn.execute("DELETE FROM room_responses WHERE room_code = ?", (room_code,))
        conn.execute("DELETE FROM room_students WHERE room_code = ?", (room_code,))
        conn.execute("DELETE FROM presentations WHERE room_code = ?", (room_code,))
        conn.execute("DELETE FROM questions WHERE room_code = ?", (room_code,))
        conn.execute(
            "DELETE FROM rooms WHERE room_code = ? AND teacher_id = ?",
            (room_code, session["teacher_id"]),
        )
        revoke_room_tab_sessions(conn, [room_code])
        conn.commit()
    except Exception:
        conn.rollback()
        conn.close()
        raise
    conn.close()
    if session.get("room_code") == room_code:
        session.pop("room_code", None)
    return jsonify({"status": "success", "message": f"Đã xóa phòng {room_code} và toàn bộ dữ liệu trong phòng."})


@app.route("/api/teacher/classes/<int:class_id>/rooms", methods=["POST"])
def create_class_room(class_id):
    if not teacher_account_required():
        return jsonify({"status": "error", "message": "Vui lòng đăng nhập tài khoản giảng viên."}), 401

    data = get_json_data()
    password = str(data.get("password", "")).strip()
    requested_room_code = str(data.get("room_code", "")).strip().upper()
    if not password:
        return jsonify({"status": "error", "message": "Mỗi phòng cần có mật khẩu."}), 400
    if requested_room_code and (
        len(requested_room_code) < 3
        or len(requested_room_code) > 32
        or any(not (character.isalnum() or character == "-") for character in requested_room_code)
    ):
        return jsonify({"status": "error", "message": "Mã phòng dài 3–32 ký tự, chỉ gồm chữ, số hoặc dấu gạch ngang."}), 400

    conn = get_db_connection()
    class_row = conn.execute(
        "SELECT id FROM classes WHERE id = ? AND teacher_id = ?",
        (class_id, session["teacher_id"]),
    ).fetchone()
    if class_row is None:
        conn.close()
        return jsonify({"status": "error", "message": "Không tìm thấy lớp của tài khoản này."}), 404

    room_code = requested_room_code or f"ARS-{secrets.token_hex(3).upper()}"
    try:
        conn.execute(
            "INSERT INTO rooms (room_code, password_hash, student_password_required, active_question_id, question_published, is_locked, teacher_id, class_id) "
            "VALUES (?, ?, ?, NULL, 0, 0, ?, ?)",
            (
                room_code,
                generate_password_hash(password),
                0,
                session["teacher_id"],
                class_id,
            ),
        )
        conn.commit()
    except sqlite3.IntegrityError:
        conn.close()
        message = "Mã phòng này đã được sử dụng. Hãy chọn mã khác." if requested_room_code else "Không thể tạo mã phòng, vui lòng thử lại."
        return jsonify({"status": "error", "message": message}), 409
    conn.close()

    session["role"] = "teacher"
    session["room_code"] = room_code
    return jsonify({"status": "success", "role": "teacher", "room_code": room_code})


def reset_room_to_idle(conn, room_code):
    """Đưa phòng về trạng thái ban đầu mỗi khi giảng viên vào phòng:
    chưa phát câu hỏi nào, nên remote/MQTT chưa được ghi nhận đáp án (ingest_mqtt_answer
    chỉ nhận khi question_published = 1 và có active_question_id).
    Chỉ đổi trạng thái phát; đáp án và bảng điểm cũ vẫn được giữ nguyên."""
    conn.execute(
        "UPDATE rooms SET active_question_id = NULL, question_published = 0, "
        "is_locked = 0, launched_at = NULL WHERE room_code = ?",
        (room_code,),
    )


@app.route("/api/teacher/rooms/<room_code>/open", methods=["POST"])
def open_class_room(room_code):
    if not teacher_account_required():
        return jsonify({"status": "error", "message": "Vui lòng đăng nhập tài khoản giảng viên."}), 401
    room_code = room_code.strip().upper()
    conn = get_db_connection()
    room = conn.execute(
        "SELECT room_code FROM rooms WHERE room_code = ? AND teacher_id = ?",
        (room_code, session["teacher_id"]),
    ).fetchone()
    if room is None:
        conn.close()
        return jsonify({"status": "error", "message": "Không tìm thấy phòng của tài khoản này."}), 404
    reset_room_to_idle(conn, room_code)
    conn.commit()
    conn.close()
    session["role"] = "teacher"
    session["room_code"] = room_code
    return jsonify({"status": "success", "role": "teacher", "room_code": room_code})


@app.route("/api/rooms/create", methods=["POST"])
def create_room():
    """Giảng viên tạo phòng mới với mật khẩu riêng."""
    if not teacher_account_required():
        return jsonify({"status": "error", "message": "Đăng nhập tài khoản giảng viên trước khi tạo phòng."}), 401
    data = get_json_data()
    password = str(data.get("password", "")).strip()
    requested_room_code = str(data.get("room_code", "")).strip().upper()
    if not password:
        return jsonify({"status": "error", "message": "Vui lòng đặt mật khẩu phòng."}), 400
    if requested_room_code and (
        len(requested_room_code) < 3
        or len(requested_room_code) > 32
        or any(not (character.isalnum() or character == "-") for character in requested_room_code)
    ):
        return jsonify({"status": "error", "message": "Mã phòng dài 3–32 ký tự, chỉ gồm chữ, số hoặc dấu gạch ngang."}), 400

    conn = get_db_connection()
    cursor = conn.cursor()
    room_code = requested_room_code or f"ARS-{secrets.token_hex(3).upper()}"
    try:
        cursor.execute(
            "INSERT INTO rooms (room_code, password_hash, student_password_required, active_question_id, question_published, is_locked, teacher_id) VALUES (?, ?, ?, NULL, 0, 0, ?)",
            (
                room_code,
                generate_password_hash(password),
                int(bool(data.get("student_password_required", True))),
                session["teacher_id"],
            ),
        )
        conn.commit()
    except sqlite3.IntegrityError:
        conn.close()
        message = "Mã phòng này đã được sử dụng. Hãy chọn mã khác." if requested_room_code else "Không thể tạo mã phòng. Vui lòng thử lại."
        return jsonify({"status": "error", "message": message}), 409
    except Exception:
        conn.close()
        raise
    conn.close()

    session["role"] = "teacher"
    session["room_code"] = room_code
    return jsonify({"status": "success", "role": "teacher", "room_code": room_code})


@app.route("/api/rooms/teacher-join", methods=["POST"])
def teacher_join_room():
    """Giảng viên vào phòng đã tồn tại bằng mã phòng và mật khẩu."""
    if not teacher_account_required():
        return jsonify({"status": "error", "message": "Đăng nhập tài khoản giảng viên trước khi vào phòng."}), 401
    data = get_json_data()
    room_code = str(data.get("room_code", "")).strip().upper()
    password = str(data.get("password", "")).strip()
    if not room_code or not password:
        return jsonify({"status": "error", "message": "Cần nhập mã phòng và mật khẩu."}), 400

    conn = get_db_connection()
    room = conn.execute(
        "SELECT room_code, password_hash, teacher_id FROM rooms WHERE room_code = ?", (room_code,)
    ).fetchone()
    conn.close()
    if room is None or not check_password_hash(room["password_hash"], password):
        return jsonify({"status": "error", "message": "Mã phòng hoặc mật khẩu không đúng."}), 401

    if room["teacher_id"] not in (None, session["teacher_id"]):
        return jsonify({"status": "error", "message": "Phòng này thuộc tài khoản giảng viên khác."}), 403

    conn = get_db_connection()
    if room["teacher_id"] is None:
        conn.execute(
            "UPDATE rooms SET teacher_id = ? WHERE room_code = ? AND teacher_id IS NULL",
            (session["teacher_id"], room_code),
        )
    reset_room_to_idle(conn, room_code)
    conn.commit()
    conn.close()

    session["role"] = "teacher"
    session["room_code"] = room_code
    return jsonify({"status": "success", "role": "teacher", "room_code": room_code})


@app.route("/api/rooms/student-join", methods=["POST"])
def student_join_room():
    """Student web login is disabled; use the controller roster mapping instead."""
    return jsonify(
        {
            "status": "error",
            "message": "Học sinh không đăng nhập web. Giảng viên nhập danh sách lớp và ID điều khiển.",
        }
    ), 410


@app.route("/api/rooms/leave", methods=["POST"])
def leave_room():
    """Kết thúc phiên tham gia phòng hiện tại."""
    if teacher_account_required():
        session.pop("room_code", None)
    else:
        session.clear()
    return jsonify({"status": "success"})


def teacher_room_required():
    if not teacher_account_required() or not session.get("room_code"):
        return False
    conn = get_db_connection()
    room = conn.execute(
        "SELECT 1 FROM rooms WHERE room_code = ? AND teacher_id = ?",
        (session["room_code"], session["teacher_id"]),
    ).fetchone()
    conn.close()
    return room is not None


def serialize_question(row):
    return {
        "id": row["id"],
        "position": row["position"] if "position" in row.keys() else None,
        "source_question_id": row["source_question_id"] if "source_question_id" in row.keys() else None,
        "text": row["text"],
        "options": json.loads(row["options_json"]),
        "correct": row["correct"],
        "explanation": row["explanation"],
    }


def serialize_presentation(row):
    return {
        "id": row["id"],
        "title": row["title"],
        "scheduled_at": row["scheduled_at"],
        "question_count": row["question_count"],
        "status": row["status"],
        "is_default": bool(row["is_default"]),
        "created_at": row["created_at"],
        "opened_at": row["opened_at"],
        "closed_at": row["closed_at"],
        "block_id": row["block_id"] or "",
        "response_count": row["response_count"] if "response_count" in row.keys() else 0,
        "responding_students": row["responding_students"] if "responding_students" in row.keys() else 0,
    }


def room_state(room_code, update_student_presence=False):
    conn = get_db_connection()
    if update_student_presence and session.get("role") == "student":
        conn.execute(
            "UPDATE room_students SET last_seen = CURRENT_TIMESTAMP WHERE room_code = ? AND student_id = ?",
            (room_code, session.get("student_id")),
        )
        conn.commit()

    room = conn.execute(
        "SELECT active_question_id, active_presentation_id, question_published, is_locked, block_id "
        "FROM rooms WHERE room_code = ?",
        (room_code,),
    ).fetchone()
    if room is None:
        conn.close()
        return None

    question_rows = conn.execute(
        "SELECT * FROM questions WHERE room_code = ? AND presentation_id = ? ORDER BY position",
        (room_code, room["active_presentation_id"]),
    ).fetchall()
    bank_question_rows = conn.execute(
        "SELECT * FROM questions WHERE room_code = ? AND presentation_id IS NULL ORDER BY position",
        (room_code,),
    ).fetchall()
    presentation = conn.execute(
        "SELECT * FROM presentations WHERE id = ? AND room_code = ?",
        (room["active_presentation_id"], room_code),
    ).fetchone()
    students = conn.execute(
        "SELECT student_id, name, roll_number, joined_at, "
        "CASE WHEN last_seen >= datetime('now', '-1 minute') THEN 1 ELSE 0 END AS online "
        "FROM room_students WHERE room_code = ? ORDER BY COALESCE(roll_number, 2147483647), student_id",
        (room_code,),
    ).fetchall()
    responses = []
    if room["active_question_id"] is not None:
        responses = conn.execute(
            "SELECT student_id, answer, timestamp FROM room_responses WHERE room_code = ? AND question_id = ?",
            (room_code, room["active_question_id"]),
        ).fetchall()
    conn.close()
    role = session.get("role")
    serialized_questions = [serialize_question(row) for row in question_rows]
    serialized_bank_questions = [serialize_question(row) for row in bank_question_rows]
    if role != "teacher" or session.get("room_code") != room_code:
        serialized_questions = [
            {key: value for key, value in question.items() if key not in {"correct", "explanation"}}
            for question in serialized_questions
        ]
        serialized_bank_questions = []
    return {
        "room_code": room_code,
        "presentation": dict(presentation) if presentation else None,
        "questions": serialized_questions,
        "question_bank": serialized_bank_questions,
        "active_question_id": room["active_question_id"],
        "question_published": bool(room["question_published"]),
        "is_locked": bool(room["is_locked"]),
        "block_id": room["block_id"] or "",
        "students": [dict(row) for row in students],
        "responses": {row["student_id"]: row["answer"] for row in responses},
    }


@app.route("/api/rooms/state", methods=["GET"])
def get_room_state():
    room_code = session.get("room_code")
    if not room_code:
        return jsonify({"status": "error", "message": "Bạn chưa vào phòng."}), 401
    if session.get("role") == "teacher" and not teacher_room_required():
        return jsonify({"status": "error", "message": "Phòng không thuộc tài khoản giảng viên này."}), 403
    if session.get("role") == "student":
        conn = get_db_connection()
        still_joined = conn.execute(
            "SELECT 1 FROM room_students WHERE room_code = ? AND student_id = ?",
            (room_code, session.get("student_id")),
        ).fetchone()
        conn.close()
        if still_joined is None:
            session.clear()
            return jsonify({"status": "error", "message": "Giảng viên đã xóa bạn khỏi phòng."}), 403
    current = room_state(room_code, update_student_presence=True)
    if current is None:
        session.clear()
        return jsonify({"status": "error", "message": "Phòng không còn tồn tại."}), 404
    return jsonify({"status": "success", "role": session.get("role"), "student_id": session.get("student_id"), "student_name": session.get("name"), **current})


@app.route("/api/rooms/gradebook", methods=["GET"])
def get_room_gradebook():
    """Trả bảng điểm mọi câu hỏi cho giảng viên trong phòng hiện tại."""
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên trong phòng mới xem được bảng điểm."}), 403

    room_code = session["room_code"]
    conn = get_db_connection()
    active_presentation = conn.execute(
        "SELECT active_presentation_id FROM rooms WHERE room_code = ?",
        (room_code,),
    ).fetchone()
    presentation_id = active_presentation["active_presentation_id"] if active_presentation else None
    questions = conn.execute(
        "SELECT id, ROW_NUMBER() OVER (ORDER BY position) AS position, text, options_json, correct, explanation "
        "FROM questions WHERE room_code = ? AND presentation_id = ? ORDER BY position",
        (room_code, presentation_id),
    ).fetchall()
    students = conn.execute(
        "SELECT student_id, name, roll_number, note, joined_at, "
        "CASE WHEN last_seen >= datetime('now', '-1 minute') THEN 1 ELSE 0 END AS online "
        "FROM room_students WHERE room_code = ? ORDER BY COALESCE(roll_number, 2147483647), student_id",
        (room_code,),
    ).fetchall()
    answers = conn.execute(
        "SELECT student_id, question_id, answer, timestamp FROM room_responses "
        "WHERE room_code = ?",
        (room_code,),
    ).fetchall()
    conn.close()

    question_data = [
        {
            "id": row["id"],
            "position": row["position"],
            "text": row["text"],
            "options": json.loads(row["options_json"]),
            "correct": row["correct"],
            "explanation": row["explanation"],
        }
        for row in questions
    ]
    question_by_id = {question["id"]: question for question in question_data}
    student_data = [dict(row) for row in students]
    response_matrix = {student["student_id"]: {} for student in student_data}
    for row in answers:
        question = question_by_id.get(row["question_id"])
        if question is not None:
            response_matrix.setdefault(row["student_id"], {})[str(row["question_id"])] = {
                "answer": row["answer"],
                "correct": row["answer"] == question["correct"],
                "timestamp": row["timestamp"],
            }

    return jsonify(
        {
            "status": "success",
            "students": student_data,
            "questions": question_data,
            "responses": response_matrix,
        }
    )


def normalize_roster_header(value):
    normalized = unicodedata.normalize("NFKD", str(value or "").casefold())
    normalized = "".join(character for character in normalized if not unicodedata.combining(character))
    normalized = normalized.replace("đ", "d").replace("_", " ")
    normalized = re.sub(r"[^a-z0-9]+", " ", normalized)
    return " ".join(normalized.split())


def parse_roster_rows(rows):
    """Parse spreadsheet rows with STT, MSSV, student name, and optional note."""
    aliases = {
        "roll_number": {"stt", "so thu tu", "thu tu", "order", "roll number", "number"},
        "student_id": {"mssv", "ma so sinh vien", "ma so sv", "student id", "student identifier"},
        "name": {"ho va ten", "ho ten", "ho ten sinh vien", "name", "student name", "full name"},
        "note": {"ghi chu", "note", "notes", "remark", "remarks"},
    }
    required_fields = {"roll_number", "student_id", "name"}
    iterator = iter(rows)
    first_rows = []
    header_row = None
    header_indices = {}
    for row_number, row in enumerate(iterator, start=1):
        first_rows.append((row_number, row))
        candidate_indices = {}
        for index, header in enumerate(row):
            normalized = normalize_roster_header(header)
            for field, names in aliases.items():
                if normalized in names:
                    candidate_indices.setdefault(field, index)
        if required_fields.issubset(candidate_indices):
            header_row = row
            header_indices = candidate_indices
            break
        if row_number >= 20:
            break
    if not first_rows:
        return [], ["Tệp danh sách đang trống."]
    if header_row is None:
        return [], ["Không tìm thấy hàng tiêu đề có đủ cột STT, MSSV, Họ và tên trong 20 hàng đầu."]

    def cell_text(value):
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    roster = []
    errors = []
    seen_ids = set()
    seen_orders = set()
    data_rows = ((row_number, row) for row_number, row in first_rows if row_number > first_rows[-1][0])
    data_rows = list(data_rows) + list(enumerate(iterator, start=first_rows[-1][0] + 1))
    for row_number, raw_row in data_rows:
        if not raw_row or not any(cell_text(value) for value in raw_row):
            continue
        fields = {
            field: cell_text(raw_row[index] if index < len(raw_row) else "")
            for field, index in header_indices.items()
        }
        try:
            roll_value = fields["roll_number"].strip()
            if re.fullmatch(r"\d+[.,]0+", roll_value):
                roll_value = re.split(r"[.,]", roll_value, maxsplit=1)[0]
            roll_number = int(roll_value.rstrip("."))
            if roll_number < 1:
                raise ValueError
        except ValueError:
            errors.append(f"Dòng {row_number}: STT phải là số nguyên dương.")
            continue
        student_id = fields["student_id"]
        name = fields["name"]
        if not student_id or not name:
            errors.append(f"Dòng {row_number}: MSSV và họ tên không được để trống.")
            continue
        if student_id in seen_ids:
            errors.append(f"Dòng {row_number}: MSSV {student_id} bị lặp.")
            continue
        if roll_number in seen_orders:
            errors.append(f"Dòng {row_number}: STT {roll_number} bị lặp.")
            continue
        seen_ids.add(student_id)
        seen_orders.add(roll_number)
        roster.append(
            {
                "roll_number": roll_number,
                "student_id": student_id,
                "name": name,
                "note": fields.get("note", ""),
                "note_provided": "note" in header_indices,
            }
        )
    if not roster and not errors:
        errors.append("Tệp không có dòng học sinh nào.")
    return roster, errors


@app.route("/api/rooms/roster/import", methods=["POST"])
def import_room_roster():
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên của phòng mới nhập được danh sách lớp."}), 403
    uploaded_file = request.files.get("file")
    if uploaded_file is None or not uploaded_file.filename:
        return jsonify({"status": "error", "message": "Chọn tệp danh sách lớp trước."}), 400
    extension = os.path.splitext(uploaded_file.filename)[1].lower()
    if extension not in {".xlsx", ".csv"}:
        return jsonify({"status": "error", "message": "Danh sách cần là tệp .xlsx hoặc .csv."}), 400
    content = uploaded_file.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024:
        return jsonify({"status": "error", "message": "Tệp danh sách tối đa 5 MB."}), 413
    try:
        if extension == ".csv":
            csv_text = content.decode("utf-8-sig")
            try:
                dialect = csv.Sniffer().sniff(csv_text[:8192], delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            roster_rows = csv.reader(StringIO(csv_text), dialect)
            roster, errors = parse_roster_rows(roster_rows)
        else:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            roster = []
            errors = ["Không tìm thấy trang tính có các cột STT, MSSV, Họ và tên."]
            for worksheet in workbook.worksheets:
                candidate_roster, candidate_errors = parse_roster_rows(
                    worksheet.iter_rows(values_only=True)
                )
                if not candidate_errors or not candidate_errors[0].startswith("Không tìm thấy hàng tiêu đề"):
                    roster, errors = candidate_roster, candidate_errors
                    break
    except (UnicodeDecodeError, InvalidFileException, BadZipFile, OSError) as error:
        return jsonify({"status": "error", "message": f"Không đọc được tệp danh sách: {error}"}), 400

    if errors:
        return jsonify({"status": "error", "message": "Danh sách chưa được nhập.", "errors": errors}), 400

    room_code = session["room_code"]
    conn = get_db_connection()
    imported_ids = {student["student_id"] for student in roster}
    existing_roster = conn.execute(
        "SELECT student_id, roll_number FROM room_students WHERE room_code = ?",
        (room_code,),
    ).fetchall()
    for student in roster:
        for existing in existing_roster:
            if existing["student_id"] in imported_ids or existing["student_id"] == student["student_id"]:
                continue
            if existing["roll_number"] == student["roll_number"]:
                conn.close()
                return jsonify(
                    {
                        "status": "error",
                        "message": f"STT {student['roll_number']} đang thuộc MSSV {existing['student_id']} chưa có trong tệp.",
                    }
                ), 409

    try:
        for student in roster:
            conn.execute(
                "INSERT INTO room_students (room_code, student_id, name, roll_number, note, last_seen) "
                "VALUES (?, ?, ?, ?, ?, '1970-01-01 00:00:00') "
                "ON CONFLICT(room_code, student_id) DO UPDATE SET "
                "name = excluded.name, roll_number = excluded.roll_number, "
                "note = CASE WHEN ? THEN excluded.note ELSE room_students.note END",
                (
                    room_code,
                    student["student_id"],
                    student["name"],
                    student["roll_number"],
                    student["note"],
                    student["note_provided"],
                ),
            )
        conn.commit()
    except Exception:
        conn.rollback()
        conn.close()
        raise
    conn.close()
    return jsonify({"status": "success", "imported": len(roster)})


@app.route("/api/rooms/roster", methods=["POST"])
def add_room_roster_student():
    """Manually add or update one student using only STT, MSSV, and name."""
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên của phòng mới quản lý được danh sách lớp."}), 403

    data = get_json_data()
    student_id = str(data.get("student_id", "")).strip()
    name = str(data.get("name", "")).strip()
    roll_value = str(data.get("roll_number", "")).strip()
    if not student_id or not name or not roll_value:
        return jsonify({"status": "error", "message": "Cần nhập đủ STT, MSSV và họ tên."}), 400
    try:
        roll_number = int(roll_value)
        if roll_number < 1:
            raise ValueError
    except ValueError:
        return jsonify({"status": "error", "message": "STT phải là số nguyên dương."}), 400

    room_code = session["room_code"]
    conn = get_db_connection()
    conflict = conn.execute(
        "SELECT student_id FROM room_students WHERE room_code = ? AND roll_number = ? AND student_id <> ?",
        (room_code, roll_number, student_id),
    ).fetchone()
    if conflict:
        conn.close()
        return jsonify({"status": "error", "message": f"STT {roll_number} đã được dùng bởi MSSV {conflict['student_id']}."}), 409

    conn.execute(
        "INSERT INTO room_students (room_code, student_id, name, roll_number, last_seen) "
        "VALUES (?, ?, ?, ?, '1970-01-01 00:00:00') "
        "ON CONFLICT(room_code, student_id) DO UPDATE SET "
        "name = excluded.name, roll_number = excluded.roll_number",
        (room_code, student_id, name, roll_number),
    )
    conn.commit()
    conn.close()
    return jsonify({"status": "success", "message": "Đã lưu học sinh vào danh sách lớp."}), 201


@app.route("/api/rooms/roster/note", methods=["POST"])
def update_room_roster_note():
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên của phòng mới sửa được ghi chú."}), 403

    data = get_json_data()
    student_id = str(data.get("student_id") or "").strip()
    note = str(data.get("note") or "").strip()
    if not student_id:
        return jsonify({"status": "error", "message": "Thiếu mã số sinh viên."}), 400
    if len(note) > 1000:
        return jsonify({"status": "error", "message": "Ghi chú tối đa 1000 ký tự."}), 400

    conn = get_db_connection()
    result = conn.execute(
        "UPDATE room_students SET note = ? WHERE room_code = ? AND student_id = ?",
        (note, session["room_code"], student_id),
    )
    updated_count = result.rowcount
    conn.commit()
    conn.close()
    if updated_count == 0:
        return jsonify({"status": "error", "message": "Không tìm thấy học sinh trong phòng."}), 404
    return jsonify({"status": "success", "message": "Đã lưu ghi chú."})

@app.route("/api/rooms/students/remove", methods=["POST"])
def remove_room_student():
    """Cho phép giảng viên xóa học sinh nhập nhầm khỏi phòng hiện tại."""
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên trong phòng mới được xóa học sinh."}), 403

    student_id = str(get_json_data().get("student_id", "")).strip()
    if not student_id:
        return jsonify({"status": "error", "message": "Thiếu mã số sinh viên cần xóa."}), 400

    room_code = session["room_code"]
    conn = get_db_connection()
    student = conn.execute(
        "SELECT 1 FROM room_students WHERE room_code = ? AND student_id = ?",
        (room_code, student_id),
    ).fetchone()
    if student is None:
        conn.close()
        return jsonify({"status": "error", "message": "Không tìm thấy học sinh trong phòng."}), 404

    conn.execute(
        "DELETE FROM room_responses WHERE room_code = ? AND student_id = ?",
        (room_code, student_id),
    )
    delete_student_exam_data(conn, room_code, student_id)
    conn.execute(
        "DELETE FROM room_students WHERE room_code = ? AND student_id = ?",
        (room_code, student_id),
    )

    # Thu hồi mọi phiên đăng nhập của MSSV này trong đúng phòng.
    student_sessions = conn.execute(
        "SELECT session_id, session_data FROM browser_sessions"
    ).fetchall()
    for browser_session in student_sessions:
        try:
            session_data = json.loads(browser_session["session_data"])
        except json.JSONDecodeError:
            continue
        if (
            session_data.get("role") == "student"
            and session_data.get("room_code") == room_code
            and session_data.get("student_id") == student_id
        ):
            conn.execute(
                "DELETE FROM browser_sessions WHERE session_id = ?",
                (browser_session["session_id"],),
            )

    conn.commit()
    conn.close()
    return jsonify({"status": "success", "message": f"Đã xóa {student_id} khỏi phòng."})


@app.route("/api/rooms/<room_code>/presentation", methods=["GET"])
def get_presentation_state(room_code):
    current = room_state(room_code.strip().upper())
    if current is None:
        return jsonify({"status": "error", "message": "Không tìm thấy phòng."}), 404
    question = next(
        (item for item in current["questions"] if item["id"] == current["active_question_id"]),
        None,
    )
    return jsonify({"status": "success", "room_code": current["room_code"], "published": current["question_published"], "question": question})


@app.route("/api/rooms/presentations", methods=["GET", "POST"])
def room_presentations():
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên trong phòng mới quản lý được bài trình chiếu."}), 403
    room_code = session["room_code"]
    conn = get_db_connection()
    bank_size = conn.execute(
        "SELECT COUNT(*) FROM questions WHERE room_code = ? AND presentation_id IS NULL",
        (room_code,),
    ).fetchone()[0]

    if request.method == "POST":
        data = get_json_data()
        title = str(data.get("title", "")).strip()[:120] or "Bài trình chiếu"
        try:
            question_count = int(data.get("question_count"))
            scheduled_at = normalize_exam_date(data.get("scheduled_at"))
        except (TypeError, ValueError) as error:
            conn.close()
            message = str(error) if isinstance(error, ValueError) else "Số câu phải là số nguyên."
            return jsonify({"status": "error", "message": message}), 400
        if bank_size == 0:
            conn.close()
            return jsonify({"status": "error", "message": "Ngân hàng câu hỏi đang trống. Hãy thêm câu hỏi trước."}), 400
        if not 1 <= question_count <= bank_size:
            conn.close()
            return jsonify({"status": "error", "message": f"Số câu phải từ 1 đến {bank_size}."}), 400
        if not scheduled_at:
            conn.close()
            return jsonify({"status": "error", "message": "Hãy chọn ngày và giờ trình chiếu."}), 400

        requested_question_ids = data.get("question_ids")
        if requested_question_ids is not None:
            if not isinstance(requested_question_ids, list) or any(
                isinstance(question_id, bool) or not isinstance(question_id, int)
                for question_id in requested_question_ids
            ):
                conn.close()
                return jsonify({"status": "error", "message": "Danh sách câu hỏi được chọn không hợp lệ."}), 400
            requested_question_ids = list(dict.fromkeys(requested_question_ids))
            if not requested_question_ids:
                conn.close()
                return jsonify({"status": "error", "message": "Hãy chọn ít nhất một câu hỏi để bốc ngẫu nhiên."}), 400
            if question_count > len(requested_question_ids):
                conn.close()
                return jsonify({
                    "status": "error",
                    "message": f"Số câu cần bốc không được vượt quá {len(requested_question_ids)} câu đã chọn.",
                }), 400
            placeholders = ",".join("?" for _ in requested_question_ids)
            bank_questions = conn.execute(
                "SELECT id, position, text, options_json, correct, explanation FROM questions "
                f"WHERE room_code = ? AND presentation_id IS NULL AND id IN ({placeholders}) ORDER BY position",
                [room_code, *requested_question_ids],
            ).fetchall()
            if len(bank_questions) != len(requested_question_ids):
                conn.close()
                return jsonify({
                    "status": "error",
                    "message": "Một hoặc nhiều câu hỏi đã chọn không còn trong ngân hàng của phòng này.",
                }), 400
        else:
            bank_questions = conn.execute(
                "SELECT id, position, text, options_json, correct, explanation FROM questions "
                "WHERE room_code = ? AND presentation_id IS NULL ORDER BY position",
                (room_code,),
            ).fetchall()
            if question_count > len(bank_questions):
                conn.close()
                return jsonify({"status": "error", "message": "Số câu vượt quá số câu có trong ngân hàng."}), 400
        chosen_questions = random.sample(bank_questions, question_count)
        cursor = conn.execute(
            "INSERT INTO presentations (room_code, title, scheduled_at, question_count, status, created_at) "
            "VALUES (?, ?, ?, ?, 'draft', ?)",
            (room_code, title, scheduled_at, question_count, time.time()),
        )
        presentation_id = cursor.lastrowid
        next_position = conn.execute(
            "SELECT COALESCE(MAX(position), 0) FROM questions WHERE room_code = ?",
            (room_code,),
        ).fetchone()[0]
        for offset, question in enumerate(chosen_questions, start=1):
            conn.execute(
                "INSERT INTO questions (room_code, position, text, options_json, correct, explanation, presentation_id, source_question_id) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    room_code,
                    next_position + offset,
                    question["text"],
                    question["options_json"],
                    question["correct"],
                    question["explanation"],
                    presentation_id,
                    question["id"],
                ),
            )
        conn.commit()
        presentation = conn.execute(
            "SELECT * FROM presentations WHERE id = ?",
            (presentation_id,),
        ).fetchone()
        conn.close()
        return jsonify({"status": "success", "presentation": serialize_presentation(presentation)}), 201

    active = conn.execute(
        "SELECT active_presentation_id FROM rooms WHERE room_code = ?",
        (room_code,),
    ).fetchone()
    rows = conn.execute(
        "SELECT p.*, COUNT(r.student_id) AS response_count, "
        "COUNT(DISTINCT r.student_id) AS responding_students "
        "FROM presentations AS p "
        "LEFT JOIN questions AS q ON q.presentation_id = p.id AND q.room_code = p.room_code "
        "LEFT JOIN room_responses AS r ON r.room_code = p.room_code AND r.question_id = q.id "
        "WHERE p.room_code = ? GROUP BY p.id ORDER BY p.is_default DESC, p.id DESC",
        (room_code,),
    ).fetchall()
    bank_questions = conn.execute(
        "SELECT id, position, text FROM questions WHERE room_code = ? "
        "AND presentation_id IS NULL ORDER BY position",
        (room_code,),
    ).fetchall()
    conn.close()
    return jsonify({
        "status": "success",
        "presentations": [serialize_presentation(row) for row in rows],
        "active_presentation_id": active["active_presentation_id"] if active else None,
        "bank_size": bank_size,
        "question_bank": [dict(row) for row in bank_questions],
    })


def get_teacher_presentation(conn, presentation_id):
    return conn.execute(
        "SELECT * FROM presentations WHERE id = ? AND room_code = ?",
        (presentation_id, session["room_code"]),
    ).fetchone()


@app.route("/api/rooms/presentations/<int:presentation_id>/status", methods=["POST"])
def set_presentation_status(presentation_id):
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên trong phòng mới điều khiển được bài trình chiếu."}), 403
    target_status = str(get_json_data().get("status", "")).strip()
    if target_status not in {"open", "closed"}:
        return jsonify({"status": "error", "message": "Trạng thái bài trình chiếu không hợp lệ."}), 400

    room_code = session["room_code"]
    conn = get_db_connection()
    presentation = get_teacher_presentation(conn, presentation_id)
    if presentation is None:
        conn.close()
        return jsonify({"status": "error", "message": "Không tìm thấy bài trình chiếu."}), 404
    if target_status == "open":
        conn.execute(
            "UPDATE presentations SET status = 'closed', closed_at = ? "
            "WHERE room_code = ? AND status = 'open' AND id != ?",
            (time.time(), room_code, presentation_id),
        )
        conn.execute(
            "UPDATE presentations SET status = 'open', opened_at = ?, closed_at = NULL WHERE id = ?",
            (time.time(), presentation_id),
        )
        conn.execute(
            "UPDATE rooms SET active_presentation_id = ?, active_question_id = NULL, "
            "question_published = 0, is_locked = 0, launched_at = NULL WHERE room_code = ?",
            (presentation_id, room_code),
        )
    elif presentation["status"] != "open":
        conn.close()
        return jsonify({"status": "error", "message": "Chỉ có thể kết thúc bài trình chiếu đang mở."}), 409
    else:
        conn.execute(
            "UPDATE presentations SET status = 'closed', closed_at = ? WHERE id = ?",
            (time.time(), presentation_id),
        )
        conn.execute(
            "UPDATE rooms SET active_question_id = NULL, question_published = 0, "
            "is_locked = 1, launched_at = NULL WHERE room_code = ? AND active_presentation_id = ?",
            (room_code, presentation_id),
        )
    conn.commit()
    conn.close()
    return jsonify({"status": "success"})


@app.route("/api/rooms/presentations/<int:presentation_id>/schedule", methods=["POST"])
def set_presentation_schedule(presentation_id):
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên trong phòng mới đổi được ngày trình chiếu."}), 403
    try:
        scheduled_at = normalize_exam_date(get_json_data().get("scheduled_at"))
    except ValueError as error:
        return jsonify({"status": "error", "message": str(error)}), 400
    conn = get_db_connection()
    if get_teacher_presentation(conn, presentation_id) is None:
        conn.close()
        return jsonify({"status": "error", "message": "Không tìm thấy bài trình chiếu."}), 404
    conn.execute("UPDATE presentations SET scheduled_at = ? WHERE id = ?", (scheduled_at, presentation_id))
    conn.commit()
    conn.close()
    return jsonify({"status": "success", "scheduled_at": scheduled_at})


@app.route("/api/rooms/presentations/<int:presentation_id>/analysis", methods=["GET"])
def get_presentation_analysis(presentation_id):
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên trong phòng mới xem được phân tích."}), 403
    conn = get_db_connection()
    presentation = get_teacher_presentation(conn, presentation_id)
    if presentation is None:
        conn.close()
        return jsonify({"status": "error", "message": "Không tìm thấy bài trình chiếu."}), 404
    questions = conn.execute(
        "SELECT id, ROW_NUMBER() OVER (ORDER BY position) AS display_position, text, options_json, correct "
        "FROM questions "
        "WHERE room_code = ? AND presentation_id = ? ORDER BY position",
        (session["room_code"], presentation_id),
    ).fetchall()
    students = conn.execute(
        "SELECT student_id, name, roll_number FROM room_students WHERE room_code = ? "
        "ORDER BY COALESCE(roll_number, 2147483647), student_id",
        (session["room_code"],),
    ).fetchall()
    answer_rows = conn.execute(
        "SELECT r.question_id, r.student_id, r.answer FROM room_responses AS r "
        "JOIN questions AS q ON q.id = r.question_id "
        "WHERE r.room_code = ? AND q.presentation_id = ?",
        (session["room_code"], presentation_id),
    ).fetchall()
    conn.close()

    answers_by_student = {}
    by_question = {}
    for answer in answer_rows:
        answers_by_student.setdefault(answer["student_id"], {})[answer["question_id"]] = answer["answer"]
        by_question.setdefault(answer["question_id"], []).append(answer["answer"])

    student_results = []
    for student in students:
        student_answers = answers_by_student.get(student["student_id"], {})
        response_map = {}
        correct_count = 0
        incorrect_count = 0
        for question in questions:
            answer = student_answers.get(question["id"])
            is_correct = answer == question["correct"] if answer is not None else None
            if is_correct is True:
                correct_count += 1
            elif is_correct is False:
                incorrect_count += 1
            response_map[str(question["id"])] = {
                "answer": answer,
                "correct": is_correct,
            }
        answered_count = correct_count + incorrect_count
        student_results.append({
            "student_id": student["student_id"],
            "name": student["name"],
            "roll_number": student["roll_number"],
            "answered_count": answered_count,
            "correct_count": correct_count,
            "incorrect_count": incorrect_count,
            "unanswered_count": len(questions) - answered_count,
            "correct_rate": round(correct_count / answered_count * 100, 1) if answered_count else None,
            "responses": response_map,
        })
    student_results.sort(
        key=lambda student: (-student["incorrect_count"], student["correct_rate"] if student["correct_rate"] is not None else 101,
                             student["roll_number"] if student["roll_number"] is not None else 2147483647,
                             student["student_id"])
    )

    analysis_questions = []
    answer_total = 0
    correct_total = 0
    for question in questions:
        answers = by_question.get(question["id"], [])
        counts = {letter: answers.count(letter) for letter in "ABCD"}
        correct_count = counts[question["correct"]]
        response_count = len(answers)
        incorrect_count = response_count - correct_count
        wrong_choices = [
            {"letter": letter, "count": counts[letter], "text": json.loads(question["options_json"]).get(letter, "")}
            for letter in "ABCD"
            if letter != question["correct"] and counts[letter]
        ]
        top_wrong_count = max((choice["count"] for choice in wrong_choices), default=0)
        answer_total += response_count
        correct_total += correct_count
        analysis_questions.append({
            "id": question["id"],
            "position": question["display_position"],
            "text": question["text"],
            "correct_answer": question["correct"],
            "response_count": response_count,
            "correct_count": correct_count,
            "incorrect_count": incorrect_count,
            "incorrect_rate": round(incorrect_count / response_count * 100, 1) if response_count else None,
            "correct_rate": round(correct_count / response_count * 100, 1) if response_count else None,
            "choice_counts": counts,
            "top_wrong": [choice for choice in wrong_choices if choice["count"] == top_wrong_count],
        })
    analysis_questions.sort(
        key=lambda question: (
            -(question["incorrect_rate"] if question["incorrect_rate"] is not None else -1),
            -question["incorrect_count"],
            question["position"],
        )
    )
    return jsonify({
        "status": "success",
        "presentation": serialize_presentation(presentation),
        "summary": {
            "class_size": len(students),
            "responding_students": sum(student["answered_count"] > 0 for student in student_results),
            "response_count": answer_total,
            "correct_rate": round(correct_total / answer_total * 100, 1) if answer_total else None,
            "incorrect_count": answer_total - correct_total,
            "question_count": len(questions),
        },
        "students": student_results,
        "questions": analysis_questions,
    })


@app.route("/api/rooms/presentations/<int:presentation_id>", methods=["DELETE"])
def delete_presentation(presentation_id):
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên trong phòng mới xóa được bài trình chiếu."}), 403
    room_code = session["room_code"]
    conn = get_db_connection()
    presentation = get_teacher_presentation(conn, presentation_id)
    if presentation is None:
        conn.close()
        return jsonify({"status": "error", "message": "Không tìm thấy bài trình chiếu."}), 404
    question_ids = [
        row["id"] for row in conn.execute(
            "SELECT id FROM questions WHERE room_code = ? AND presentation_id = ?",
            (room_code, presentation_id),
        )
    ]
    if question_ids:
        placeholders = ",".join("?" for _ in question_ids)
        conn.execute(
            f"DELETE FROM room_responses WHERE room_code = ? AND question_id IN ({placeholders})",
            [room_code, *question_ids],
        )
        conn.execute(f"DELETE FROM questions WHERE id IN ({placeholders})", question_ids)
    active = conn.execute(
        "SELECT active_presentation_id FROM rooms WHERE room_code = ?",
        (room_code,),
    ).fetchone()
    if active and active["active_presentation_id"] == presentation_id:
        fallback_presentation = conn.execute(
            "SELECT id FROM presentations WHERE room_code = ? AND id != ? "
            "ORDER BY is_default DESC, id DESC LIMIT 1",
            (room_code, presentation_id),
        ).fetchone()
        if fallback_presentation:
            conn.execute(
                "UPDATE presentations SET status = 'open', opened_at = ?, closed_at = NULL WHERE id = ?",
                (time.time(), fallback_presentation["id"]),
            )
        conn.execute(
            "UPDATE rooms SET active_presentation_id = ?, active_question_id = NULL, "
            "question_published = 0, is_locked = ?, launched_at = NULL WHERE room_code = ?",
            (fallback_presentation["id"] if fallback_presentation else None,
             int(fallback_presentation is None), room_code),
        )
    conn.execute("DELETE FROM presentations WHERE id = ?", (presentation_id,))
    conn.execute(
        "UPDATE rooms SET presentations_initialized = 1 WHERE room_code = ?",
        (room_code,),
    )
    conn.commit()
    conn.close()
    return jsonify({"status": "success"})


@app.route("/api/rooms/questions", methods=["POST"])
def add_room_question():
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên trong phòng mới được thêm câu hỏi."}), 403
    data = get_json_data()
    question_text = str(data.get("text", "")).strip()
    options = data.get("options")
    correct = str(data.get("correct", "")).strip().upper()
    if not question_text or not isinstance(options, dict) or any(not str(options.get(key, "")).strip() for key in "ABCD") or correct not in {"A", "B", "C", "D"}:
        return jsonify({"status": "error", "message": "Cần nhập câu hỏi, đủ 4 đáp án và đáp án đúng A/B/C/D."}), 400

    room_code = session["room_code"]
    conn = get_db_connection()
    position = conn.execute(
        "SELECT COALESCE(MAX(position), 0) + 1 FROM questions WHERE room_code = ?", (room_code,)
    ).fetchone()[0]
    cursor = conn.execute(
        "INSERT INTO questions (room_code, position, text, options_json, correct, explanation, presentation_id) "
        "VALUES (?, ?, ?, ?, ?, ?, NULL)",
        (room_code, position, question_text, json.dumps({key: str(options[key]).strip() for key in "ABCD"}, ensure_ascii=False), correct, str(data.get("explanation", "")).strip()),
    )
    source_question = conn.execute("SELECT * FROM questions WHERE id = ?", (cursor.lastrowid,)).fetchone()
    default_presentation = conn.execute(
        "SELECT id FROM presentations WHERE room_code = ? AND is_default = 1",
        (room_code,),
    ).fetchone()
    if default_presentation:
        presentation_position = conn.execute(
            "SELECT COALESCE(MAX(position), 0) + 1 FROM questions WHERE room_code = ?",
            (room_code,),
        ).fetchone()[0]
        conn.execute(
            "INSERT INTO questions (room_code, position, text, options_json, correct, explanation, presentation_id, source_question_id) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            (
                room_code,
                presentation_position,
                source_question["text"],
                source_question["options_json"],
                source_question["correct"],
                source_question["explanation"],
                default_presentation["id"],
                source_question["id"],
            ),
        )
        conn.execute(
            "UPDATE presentations SET question_count = question_count + 1 WHERE id = ?",
            (default_presentation["id"],),
        )
    conn.commit()
    conn.close()
    return jsonify({"status": "success", "question": serialize_question(source_question)}), 201


def parse_question_text_file(text):
    """Parse labeled question blocks, with or without blank lines between them."""

    def normalize_label(label):
        normalized = unicodedata.normalize("NFKD", label.casefold())
        normalized = "".join(character for character in normalized if not unicodedata.combining(character))
        return " ".join(normalized.replace("đ", "d").split())

    def classify_line(line):
        stripped = line.strip()
        if not stripped:
            return None, ""

        option_match = re.match(r"^([A-D])\s*[.):\-]\s*(.*)$", stripped, re.IGNORECASE)
        if option_match:
            return option_match.group(1).upper(), option_match.group(2).strip()

        correct_match = re.match(
            r"^(?:đáp\s*án(?:\s*đúng)?|dap\s*an(?:\s*dung)?|correct(?:\s+answer)?)\s*(?:là\s*)?[:=\-]?\s*([A-D])(?:\b|\s|$)",
            stripped,
            re.IGNORECASE,
        )
        if correct_match:
            return "correct", correct_match.group(1).upper()

        explanation_match = re.match(r"^(?:giải\s*thích|giai\s*thich|explanation)\s*[:=]\s*(.*)$", stripped, re.IGNORECASE)
        if explanation_match:
            return "explanation", explanation_match.group(1).strip()

        question_match = re.match(
            r"^(?:câu\s*hỏi(?:\s*\d+)?|cau\s*hoi(?:\s*\d+)?|question(?:\s*\d+)?|q\s*\d*|câu\s*\d+|cau\s*\d+)\s*[:.)\-]\s*(.*)$",
            stripped,
            re.IGNORECASE,
        )
        if question_match:
            return "text", question_match.group(1).strip()

        # Also accept numbered question headings such as "1. What is ...?".
        numbered_question = re.match(r"^\d+[.)]\s+(.+)$", stripped)
        if numbered_question:
            return "text", numbered_question.group(1).strip()

        return "continuation", stripped

    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks = []
    current_block = []
    for line in lines:
        field, _ = classify_line(line)
        if field == "text" and current_block:
            if any(classify_line(existing)[0] == "text" for existing in current_block):
                blocks.append(current_block)
                current_block = []
        if not line.strip():
            if current_block:
                blocks.append(current_block)
                current_block = []
        else:
            current_block.append(line)
    if current_block:
        blocks.append(current_block)

    questions = []
    errors = []
    for block_index, block in enumerate(blocks, start=1):
        fields = {"options": {}}
        current_field = None
        for line in block:
            field, value = classify_line(line)
            if field is None:
                continue
            if field == "continuation":
                if current_field in {"A", "B", "C", "D"}:
                    fields["options"][current_field] = f"{fields['options'].get(current_field, '')}\n{value}".strip()
                elif current_field:
                    fields[current_field] = f"{fields.get(current_field, '')}\n{value}".strip()
                elif not fields.get("text"):
                    fields["text"] = value
                    current_field = "text"
                else:
                    errors.append(f"Câu {block_index}: không nhận ra dòng '{value}'.")
                continue
            if field in {"A", "B", "C", "D"}:
                fields["options"][field] = value
            else:
                fields[field] = value
            current_field = field

        correct = str(fields.get("correct", "")).strip().upper()
        if not fields.get("text"):
            errors.append(f"Câu {block_index}: thiếu nội dung câu hỏi.")
        missing_options = [letter for letter in "ABCD" if not fields["options"].get(letter, "").strip()]
        if missing_options:
            errors.append(f"Câu {block_index}: thiếu đáp án {', '.join(missing_options)}.")
        if correct not in {"A", "B", "C", "D"}:
            errors.append(f"Câu {block_index}: đáp án đúng phải là A, B, C hoặc D.")
        if fields.get("text") and not missing_options and correct in {"A", "B", "C", "D"}:
            questions.append(
                {
                    "text": fields["text"].strip(),
                    "options": {letter: fields["options"][letter].strip() for letter in "ABCD"},
                    "correct": correct,
                    "explanation": str(fields.get("explanation", "")).strip(),
                }
            )

    if not blocks:
        errors.append("Tệp chưa có nội dung câu hỏi.")
    return questions, errors


@app.route("/api/rooms/questions/import", methods=["POST"])
def import_room_questions():
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên của phòng mới nhập được câu hỏi."}), 403
    uploaded_file = request.files.get("file")
    if uploaded_file is None or not uploaded_file.filename:
        return jsonify({"status": "error", "message": "Chọn tệp câu hỏi .txt trước."}), 400
    if not uploaded_file.filename.lower().endswith((".txt", ".text")):
        return jsonify({"status": "error", "message": "Chỉ hỗ trợ tệp .txt hoặc .text."}), 400
    raw_content = uploaded_file.read(1024 * 1024 + 1)
    if len(raw_content) > 1024 * 1024:
        return jsonify({"status": "error", "message": "Tệp tối đa 1 MB."}), 413
    try:
        text = raw_content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return jsonify({"status": "error", "message": "Tệp phải được lưu dưới dạng UTF-8."}), 400

    questions, errors = parse_question_text_file(text)
    if errors:
        return jsonify({"status": "error", "message": "Tệp có lỗi; chưa câu hỏi nào được nhập.", "errors": errors}), 400

    room_code = session["room_code"]
    conn = get_db_connection()
    position = conn.execute(
        "SELECT COALESCE(MAX(position), 0) FROM questions WHERE room_code = ?", (room_code,)
    ).fetchone()[0]
    default_presentation = conn.execute(
        "SELECT id FROM presentations WHERE room_code = ? AND is_default = 1",
        (room_code,),
    ).fetchone()
    presentation_position = conn.execute(
        "SELECT COALESCE(MAX(position), 0) FROM questions WHERE room_code = ?",
        (room_code,),
    ).fetchone()[0]
    try:
        for offset, question in enumerate(questions, start=1):
            conn.execute(
                "INSERT INTO questions (room_code, position, text, options_json, correct, explanation) "
                "VALUES (?, ?, ?, ?, ?, ?)",
                (
                    room_code,
                    position + offset,
                    question["text"],
                    json.dumps(question["options"], ensure_ascii=False),
                    question["correct"],
                    question["explanation"],
                ),
            )
            if default_presentation:
                conn.execute(
                    "INSERT INTO questions (room_code, position, text, options_json, correct, explanation, presentation_id, source_question_id) "
                    "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                    (
                        room_code,
                        presentation_position + len(questions) + offset,
                        question["text"],
                        json.dumps(question["options"], ensure_ascii=False),
                        question["correct"],
                        question["explanation"],
                        default_presentation["id"],
                        conn.execute(
                            "SELECT id FROM questions WHERE room_code = ? AND position = ? AND presentation_id IS NULL",
                            (room_code, position + offset),
                        ).fetchone()["id"],
                    ),
                )
        if default_presentation:
            conn.execute(
                "UPDATE presentations SET question_count = question_count + ? WHERE id = ?",
                (len(questions), default_presentation["id"]),
            )
        conn.commit()
    except Exception:
        conn.rollback()
        conn.close()
        raise
    conn.close()
    return jsonify({"status": "success", "imported": len(questions)})


@app.route("/api/rooms/roster.xlsx", methods=["GET"])
@app.route("/api/rooms/gradebook.xlsx", methods=["GET"])
def export_room_roster():
    """Xuất danh sách lớp gồm STT, họ tên, MSSV và ghi chú."""
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên của phòng mới xuất được danh sách lớp."}), 403

    room_code = session["room_code"]
    conn = get_db_connection()
    students = conn.execute(
        "SELECT roll_number, name, student_id, note FROM room_students "
        "WHERE room_code = ? ORDER BY COALESCE(roll_number, 2147483647), student_id",
        (room_code,),
    ).fetchall()
    conn.close()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Danh sách lớp"
    headers = ["STT", "Họ và tên", "MSSV", "Ghi chú"]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="1E3A8A")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def spreadsheet_text(value):
        value = "" if value is None else str(value)
        if value.startswith(("=", "+", "-", "@")):
            return "'" + value
        return value

    for student in students:
        sheet.append([
            student["roll_number"] or "",
            spreadsheet_text(student["name"]),
            spreadsheet_text(student["student_id"]),
            spreadsheet_text(student["note"]),
        ])

    sheet.column_dimensions["A"].width = 8
    sheet.column_dimensions["B"].width = 28
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 40
    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name=f"Danh_Sach_Lop_{room_code}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/api/rooms/questions/<int:question_id>/launch", methods=["POST"])
def launch_room_question(question_id):
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên trong phòng mới được phát câu hỏi."}), 403
    conn = get_db_connection()
    active_presentation = conn.execute(
        "SELECT r.active_presentation_id, p.status FROM rooms AS r "
        "LEFT JOIN presentations AS p ON p.id = r.active_presentation_id "
        "WHERE r.room_code = ?",
        (session["room_code"],),
    ).fetchone()
    if active_presentation is None or active_presentation["status"] != "open":
        conn.close()
        return jsonify({"status": "error", "message": "Hãy mở bài trình chiếu trước khi phát câu hỏi."}), 409
    question = conn.execute(
        "SELECT id FROM questions WHERE id = ? AND room_code = ? AND presentation_id = ?",
        (
            question_id,
            session["room_code"],
            active_presentation["active_presentation_id"] if active_presentation else None,
        ),
    ).fetchone()
    if question is None:
        conn.close()
        return jsonify({"status": "error", "message": "Không tìm thấy câu hỏi."}), 404
    conn.execute(
        "UPDATE rooms SET active_question_id = ?, question_published = 1, is_locked = 0, launched_at = ? WHERE room_code = ?",
        (question_id, time.time(), session["room_code"]),
    )
    conn.commit()
    conn.close()
    return jsonify({"status": "success"})


@app.route("/api/rooms/next-question", methods=["POST"])
def next_room_question():
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên mới được chuyển câu hỏi."}), 403
    room_code = session["room_code"]
    conn = get_db_connection()
    current = conn.execute(
        "SELECT r.active_question_id, r.active_presentation_id, p.status "
        "FROM rooms AS r LEFT JOIN presentations AS p ON p.id = r.active_presentation_id "
        "WHERE r.room_code = ?",
        (room_code,),
    ).fetchone()
    if current is None or current["status"] != "open":
        conn.close()
        return jsonify({"status": "error", "message": "Hãy mở bài trình chiếu trước khi chuyển câu hỏi."}), 409
    if current["active_question_id"] is None:
        conn.close()
        return jsonify({"status": "error", "message": "Chưa có câu hỏi nào được phát."}), 400
    current_position = conn.execute(
        "SELECT position FROM questions WHERE id = ? AND room_code = ?",
        (current["active_question_id"], room_code),
    ).fetchone()
    next_question = conn.execute(
        "SELECT id FROM questions WHERE room_code = ? AND presentation_id = ? AND position > ? "
        "ORDER BY position LIMIT 1",
        (room_code, current["active_presentation_id"], current_position["position"]),
    ).fetchone()
    if next_question is None:
        conn.close()
        return jsonify({"status": "error", "message": "Đây đã là câu hỏi cuối cùng."}), 400
    conn.execute(
        "UPDATE rooms SET active_question_id = ?, question_published = 1, is_locked = 0, launched_at = ? WHERE room_code = ?",
        (next_question["id"], time.time(), room_code),
    )
    conn.commit()
    conn.close()
    return jsonify({"status": "success"})


@app.route("/api/rooms/lock", methods=["POST"])
def toggle_room_lock():
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên mới được khóa câu hỏi."}), 403
    data = get_json_data()
    conn = get_db_connection()
    active_presentation = conn.execute(
        "SELECT p.status FROM rooms AS r LEFT JOIN presentations AS p "
        "ON p.id = r.active_presentation_id WHERE r.room_code = ?",
        (session["room_code"],),
    ).fetchone()
    if active_presentation is None or active_presentation["status"] != "open":
        conn.close()
        return jsonify({"status": "error", "message": "Hãy mở bài trình chiếu trước khi khóa hoặc mở nhận đáp án."}), 409
    conn.execute(
        "UPDATE rooms SET is_locked = ? WHERE room_code = ?",
        (int(bool(data.get("locked"))), session["room_code"]),
    )
    conn.commit()
    conn.close()
    return jsonify({"status": "success"})


@app.route("/api/rooms/presentations/<int:presentation_id>/block", methods=["POST"])
def set_presentation_block(presentation_id):
    """Gán một khối remote riêng cho bài trình chiếu."""
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên trong phòng mới được gán khối remote."}), 403
    raw = str(get_json_data().get("block_id", "")).strip()
    conn = get_db_connection()
    presentation = get_teacher_presentation(conn, presentation_id)
    if presentation is None:
        conn.close()
        return jsonify({"status": "error", "message": "Không tìm thấy bài trình chiếu."}), 404
    if not raw:
        conn.execute("UPDATE presentations SET block_id = NULL WHERE id = ?", (presentation_id,))
        conn.commit()
        conn.close()
        return jsonify({"status": "success", "block_id": "", "message": "Đã bỏ gán khối remote."})
    block_id = normalize_block_id(raw)
    if not block_id:
        conn.close()
        return jsonify({"status": "error", "message": "Khối ID chỉ gồm chữ, số, - hoặc _ (tối đa 32 ký tự)."}), 400
    # Một khối chỉ được gán cho một bài trình chiếu tại một thời điểm.
    released = [
        row["title"]
        for row in conn.execute(
            "SELECT title FROM presentations WHERE UPPER(block_id) = ? AND id <> ?",
            (block_id, presentation_id),
        )
    ]
    conn.execute("UPDATE presentations SET block_id = NULL WHERE UPPER(block_id) = ? AND id <> ?", (block_id, presentation_id))
    conn.execute("UPDATE presentations SET block_id = ? WHERE id = ?", (block_id, presentation_id))
    conn.commit()
    conn.close()
    message = f"Bài trình chiếu đang nhận đáp án từ khối {block_id}."
    if released:
        message += f" (Đã nhả khối này khỏi bài {', '.join(released)}.)"
    return jsonify({
        "status": "success",
        "block_id": block_id,
        "message": message,
        "released_presentations": released,
    })


@app.route("/api/rooms/answer", methods=["POST"])
def submit_room_answer():
    return jsonify(
        {
            "status": "error",
            "message": "Đáp án được nhận từ bộ điều khiển qua MQTT.",
        }
    ), 410


@app.route("/api/responses", methods=["GET"])
def get_responses():
    """Lấy toàn bộ phản hồi để hiển thị biểu đồ realtime"""
    question_id = request.args.get("question_id")
    conn = get_db_connection()
    cursor = conn.cursor()

    if question_id:
        cursor.execute(
            "SELECT * FROM responses WHERE question_id = ? ORDER BY timestamp DESC",
            (question_id,),
        )
    else:
        cursor.execute("SELECT * FROM responses ORDER BY timestamp DESC")

    rows = cursor.fetchall()
    conn.close()

    responses = [dict(row) for row in rows]
    return jsonify({"status": "success", "data": responses})


@app.route("/api/submit", methods=["POST"])
def submit_answer():
    """API nhận đáp án từ Web Clicker hoặc Mô phỏng N-thiết bị"""
    data = request.json or {}
    student_id = data.get("student_id")
    question_id = data.get("question_id")
    answer = data.get("answer")

    if not student_id or not question_id or not answer:
        return (
            jsonify({"status": "error", "message": "Thiếu dữ liệu đầu vào"}),
            400,
        )

    save_response_to_db(student_id, question_id, answer)
    return jsonify({"status": "success", "message": "Đã ghi nhận câu trả lời"})


@app.route("/api/clear", methods=["POST"])
def clear_responses():
    """Xóa dữ liệu thống kê để làm mới câu hỏi"""
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM responses")
    conn.commit()
    conn.close()
    return jsonify(
        {"status": "success", "message": "Đã xóa toàn bộ dữ liệu phản hồi"}
    )


# -----------------------------------------------------------------------------
# Kiểm tra trực tiếp trên app
#   - Độc lập với màn hình trình chiếu và remote/MQTT.
#   - Giảng viên đặt thời gian và số câu; mỗi sinh viên nhận bộ câu hỏi bốc ngẫu
#     nhiên từ ngân hàng câu hỏi của phòng, thứ tự câu và thứ tự đáp án xáo riêng.
#   - Đồng hồ và chấm điểm do máy chủ quyết định (không tin đồng hồ của máy sinh viên).
# -----------------------------------------------------------------------------
EXAM_MAX_MINUTES = 300
EXAM_CODE_ALPHABET = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"  # bỏ I, O, 0, 1 cho dễ đọc
EXAM_LETTERS = ("A", "B", "C", "D")
_exam_rng = random.SystemRandom()


def generate_exam_code(conn):
    for _ in range(30):
        code = "".join(secrets.choice(EXAM_CODE_ALPHABET) for _ in range(6))
        if conn.execute("SELECT 1 FROM exams WHERE access_code = ?", (code,)).fetchone() is None:
            return code
    raise RuntimeError("Không tạo được mã bài kiểm tra duy nhất.")


def exam_score10(correct, total):
    return round(correct / total * 10, 2) if total else 0


EXAM_DATE_INPUT_FORMATS = ("%Y-%m-%dT%H:%M", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d")


def normalize_exam_date(value):
    """Chuẩn hóa ngày kiểm tra giảng viên nhập thành 'YYYY-MM-DDTHH:MM' (giờ địa phương, không múi giờ).

    Để trống -> None. Chỉ nhập ngày thì lấy 00:00. Sai định dạng -> ValueError.
    """
    text = str(value or "").strip()
    if not text:
        return None
    for fmt in EXAM_DATE_INPUT_FORMATS:
        try:
            parsed = datetime.strptime(text, fmt)
        except ValueError:
            continue
        if not 2000 <= parsed.year <= 2100:
            break
        return parsed.strftime("%Y-%m-%dT%H:%M")
    raise ValueError("Ngày kiểm tra không hợp lệ. Hãy chọn lại ngày và giờ.")


def format_exam_date(value):
    """'2026-10-05T07:30' -> '05/10/2026 07:30' (chuỗi rỗng nếu chưa đặt ngày)."""
    if not value:
        return ""
    try:
        return datetime.strptime(value, "%Y-%m-%dT%H:%M").strftime("%d/%m/%Y %H:%M")
    except ValueError:
        return str(value)


def serialize_exam(row):
    return {
        "scheduled_at": row["scheduled_at"],
        "id": row["id"],
        "title": row["title"],
        "duration_minutes": row["duration_minutes"],
        "question_count": row["question_count"],
        "access_code": row["access_code"],
        "status": row["status"],
        "show_answers": bool(row["show_answers"]),
        "created_at": row["created_at"],
        "opened_at": row["opened_at"],
        "closed_at": row["closed_at"],
    }


def get_exam_question_ids(conn, exam):
    """Return the exam's configured source-question pool, or the full bank for legacy exams."""
    configured_ids = json.loads(exam["question_pool_json"]) if exam["question_pool_json"] else None
    query = "SELECT id FROM questions WHERE room_code = ? AND presentation_id IS NULL"
    parameters = [exam["room_code"]]
    if configured_ids is not None:
        if not configured_ids:
            return []
        placeholders = ",".join("?" for _ in configured_ids)
        query += f" AND id IN ({placeholders})"
        parameters.extend(configured_ids)
    return [row["id"] for row in conn.execute(query, parameters)]


def delete_exam_data_for_rooms(conn, room_codes):
    """Xóa bài kiểm tra, lượt làm bài và đáp án thuộc các phòng (dùng khi xóa phòng/lớp)."""
    room_codes = list(room_codes)
    if not room_codes:
        return
    placeholders = ",".join("?" for _ in room_codes)
    exam_ids = f"SELECT id FROM exams WHERE room_code IN ({placeholders})"
    conn.execute(
        f"DELETE FROM exam_answers WHERE attempt_id IN "
        f"(SELECT id FROM exam_attempts WHERE exam_id IN ({exam_ids}))",
        room_codes,
    )
    conn.execute(f"DELETE FROM exam_attempts WHERE exam_id IN ({exam_ids})", room_codes)
    conn.execute(f"DELETE FROM exams WHERE room_code IN ({placeholders})", room_codes)


def delete_student_exam_data(conn, room_code, student_id):
    """Xóa các lượt làm bài của một sinh viên trong phòng (dùng khi xóa sinh viên khỏi lớp)."""
    conn.execute(
        "DELETE FROM exam_answers WHERE attempt_id IN (SELECT a.id FROM exam_attempts AS a "
        "JOIN exams AS e ON e.id = a.exam_id WHERE e.room_code = ? AND a.student_id = ?)",
        (room_code, student_id),
    )
    conn.execute(
        "DELETE FROM exam_attempts WHERE student_id = ? AND exam_id IN "
        "(SELECT id FROM exams WHERE room_code = ?)",
        (student_id, room_code),
    )


def grade_attempt(conn, attempt):
    """Chấm điểm theo đáp án đã lưu. Trả về (số câu đúng, tổng số câu)."""
    items = json.loads(attempt["items_json"])
    question_ids = [item["qid"] for item in items]
    if not question_ids:
        return 0, 0
    placeholders = ",".join("?" for _ in question_ids)
    correct_by_question = {
        row["id"]: row["correct"]
        for row in conn.execute(
            f"SELECT id, correct FROM questions WHERE id IN ({placeholders})", question_ids
        )
    }
    chosen_by_question = {
        row["question_id"]: row["choice"]
        for row in conn.execute(
            "SELECT question_id, choice FROM exam_answers WHERE attempt_id = ?", (attempt["id"],)
        )
    }
    correct = 0
    for item in items:
        shown_letter = chosen_by_question.get(item["qid"])
        if shown_letter in EXAM_LETTERS:
            original_letter = item["order"][EXAM_LETTERS.index(shown_letter)]
            if original_letter == correct_by_question.get(item["qid"]):
                correct += 1
    return correct, len(items)


def finalize_attempt(conn, attempt, reason):
    """Chốt bài làm (chỉ chạy một lần). reason: manual | timeout | closed."""
    if attempt["submitted_at"] is not None:
        return
    correct, total = grade_attempt(conn, attempt)
    # Hết giờ thì ghi nhận thời điểm nộp đúng bằng hạn chót, không phải lúc máy chủ phát hiện.
    submitted_at = attempt["deadline"] if reason == "timeout" else time.time()
    conn.execute(
        "UPDATE exam_attempts SET submitted_at = ?, submit_reason = ?, correct_count = ?, total = ? "
        "WHERE id = ? AND submitted_at IS NULL",
        (submitted_at, reason, correct, total, attempt["id"]),
    )


def refresh_attempt(conn, attempt_id):
    return conn.execute("SELECT * FROM exam_attempts WHERE id = ?", (attempt_id,)).fetchone()


def auto_finalize(conn, attempt, exam):
    """Tự nộp khi hết giờ hoặc khi giảng viên đã kết thúc bài. Trả về bản ghi mới nhất."""
    if attempt["submitted_at"] is not None:
        return attempt
    if time.time() >= attempt["deadline"]:
        finalize_attempt(conn, attempt, "timeout")
    elif exam["status"] == "closed":
        finalize_attempt(conn, attempt, "closed")
    else:
        return attempt
    conn.commit()
    return refresh_attempt(conn, attempt["id"])


def finalize_due_attempts(conn, exam):
    """Chốt mọi bài đã hết giờ (hoặc tất cả nếu bài đã kết thúc) để giảng viên thấy kết quả đúng."""
    if exam["status"] == "closed":
        rows = conn.execute(
            "SELECT * FROM exam_attempts WHERE exam_id = ? AND submitted_at IS NULL", (exam["id"],)
        ).fetchall()
    else:
        rows = conn.execute(
            "SELECT * FROM exam_attempts WHERE exam_id = ? AND submitted_at IS NULL AND deadline <= ?",
            (exam["id"], time.time()),
        ).fetchall()
    for row in rows:
        finalize_attempt(conn, row, "timeout" if time.time() >= row["deadline"] else "closed")
    if rows:
        conn.commit()


def clean_student_id(value):
    return str(value or "").strip()


def build_attempt_payload(conn, attempt, exam, include_token=False):
    """Dữ liệu gửi cho sinh viên. Tuyệt đối không chứa đáp án đúng khi đang làm bài."""
    student = conn.execute(
        "SELECT name FROM room_students WHERE room_code = ? AND student_id = ?",
        (exam["room_code"], attempt["student_id"]),
    ).fetchone()
    items = json.loads(attempt["items_json"])
    question_rows = {}
    if items:
        placeholders = ",".join("?" for _ in items)
        question_rows = {
            row["id"]: row
            for row in conn.execute(
                f"SELECT * FROM questions WHERE id IN ({placeholders})", [item["qid"] for item in items]
            )
        }
    answers = {
        str(row["question_id"]): row["choice"]
        for row in conn.execute(
            "SELECT question_id, choice FROM exam_answers WHERE attempt_id = ?", (attempt["id"],)
        )
    }

    def shown_question(index, item):
        row = question_rows.get(item["qid"])
        if row is None:
            return None
        options = json.loads(row["options_json"])
        return {
            "index": index,
            "qid": row["id"],
            "text": row["text"],
            "options": [
                {"letter": EXAM_LETTERS[position], "text": options[original]}
                for position, original in enumerate(item["order"])
            ],
        }, row

    submitted = attempt["submitted_at"] is not None
    payload = {
        "status": "success",
        "exam": {
            "title": exam["title"],
            "duration_minutes": exam["duration_minutes"],
            "status": exam["status"],
        },
        "student": {"student_id": attempt["student_id"], "name": student["name"] if student else ""},
        "submitted": submitted,
        "total": len(items),
        "answers": answers,
    }
    if include_token:
        payload["token"] = attempt["token"]

    if not submitted:
        payload["remaining_seconds"] = max(0, int(attempt["deadline"] - time.time()))
        payload["questions"] = [
            built[0]
            for built in (shown_question(index, item) for index, item in enumerate(items, start=1))
            if built is not None
        ]
        return payload

    correct = attempt["correct_count"] or 0
    total = attempt["total"] or len(items)
    payload["result"] = {
        "correct": correct,
        "total": total,
        "score10": exam_score10(correct, total),
        "reason": attempt["submit_reason"],
        "submitted_at": attempt["submitted_at"],
        "seconds_used": max(0, int(attempt["submitted_at"] - attempt["started_at"])),
    }
    if exam["show_answers"]:
        review = []
        for index, item in enumerate(items, start=1):
            built = shown_question(index, item)
            if built is None:
                continue
            question, row = built
            correct_letter = EXAM_LETTERS[item["order"].index(row["correct"])]
            chosen = answers.get(str(row["id"]))
            question.update(
                {
                    "chosen": chosen,
                    "correct_letter": correct_letter,
                    "is_correct": chosen == correct_letter,
                    "explanation": row["explanation"],
                }
            )
            review.append(question)
        payload["review"] = review
    return payload


def load_attempt_from_token(conn):
    """Nạp lượt làm bài theo header X-Attempt-Token. Trả về (attempt, exam) hoặc (None, None)."""
    token = request.headers.get("X-Attempt-Token", "").strip()
    if not token:
        return None, None
    attempt = conn.execute("SELECT * FROM exam_attempts WHERE token = ?", (token,)).fetchone()
    if attempt is None:
        return None, None
    exam = conn.execute("SELECT * FROM exams WHERE id = ?", (attempt["exam_id"],)).fetchone()
    if exam is None:
        return None, None
    return attempt, exam


def invalid_attempt_response():
    return jsonify(
        {
            "status": "error",
            "code": "invalid_session",
            "message": "Phiên làm bài không hợp lệ hoặc đã bị xóa. Hãy vào lại bằng mã bài và MSSV.",
        }
    ), 401


# ----- API giảng viên -----------------------------------------------------------
def get_teacher_exam(conn, exam_id):
    return conn.execute(
        "SELECT * FROM exams WHERE id = ? AND room_code = ?", (exam_id, session["room_code"])
    ).fetchone()


def exam_not_found_response():
    return jsonify({"status": "error", "message": "Không tìm thấy bài kiểm tra trong phòng này."}), 404


def build_exam_results(conn, exam):
    """Ma trận kết quả giống bảng điểm remote: mỗi sinh viên một dòng, mỗi câu hỏi một cột.

    Mỗi sinh viên nhận bộ câu khác nhau nên các cột là hợp của mọi câu đã được bốc ra đề.
    Ô của một câu không nằm trong đề của sinh viên đó sẽ không có mục trong "responses".
    Đáp án hiển thị là chữ cái gốc trong ngân hàng câu hỏi (không phải chữ cái đã bị xáo).
    Trả về (rows, summary, questions).
    """
    finalize_due_attempts(conn, exam)
    students = conn.execute(
        "SELECT student_id, name, roll_number FROM room_students WHERE room_code = ? "
        "ORDER BY COALESCE(roll_number, 2147483647), student_id",
        (exam["room_code"],),
    ).fetchall()
    attempts = {
        row["student_id"]: row
        for row in conn.execute("SELECT * FROM exam_attempts WHERE exam_id = ?", (exam["id"],))
    }
    answers_by_attempt = {}
    for row in conn.execute(
        "SELECT attempt_id, question_id, choice FROM exam_answers WHERE attempt_id IN "
        "(SELECT id FROM exam_attempts WHERE exam_id = ?)",
        (exam["id"],),
    ):
        answers_by_attempt.setdefault(row["attempt_id"], {})[row["question_id"]] = row["choice"]

    assigned_ids = set()
    for attempt in attempts.values():
        assigned_ids.update(item["qid"] for item in json.loads(attempt["items_json"]))
    question_rows = {}
    if assigned_ids:
        placeholders = ",".join("?" for _ in assigned_ids)
        question_rows = {
            row["id"]: row
            for row in conn.execute(
                f"SELECT id, position, text, options_json, correct, explanation FROM questions "
                f"WHERE id IN ({placeholders}) ORDER BY position",
                list(assigned_ids),
            )
        }
    questions = [
        {
            "id": row["id"],
            "position": row["position"],
            "text": row["text"],
            "options": json.loads(row["options_json"]),
            "correct": row["correct"],
            "explanation": row["explanation"],
        }
        for row in question_rows.values()
    ]

    rows = []
    summary = {"class_size": len(students), "not_started": 0, "in_progress": 0, "submitted": 0}
    score_total = 0.0
    for student in students:
        attempt = attempts.get(student["student_id"])
        entry = {
            "student_id": student["student_id"],
            "name": student["name"],
            "roll_number": student["roll_number"],
            "state": "not_started",
            "responses": {},
        }
        if attempt is not None:
            items = json.loads(attempt["items_json"])
            answers = answers_by_attempt.get(attempt["id"], {})
            responses = {}
            for item in items:
                question = question_rows.get(item["qid"])
                if question is None:
                    continue
                shown = answers.get(item["qid"])
                if shown in EXAM_LETTERS:
                    original = item["order"][EXAM_LETTERS.index(shown)]
                    responses[str(item["qid"])] = {
                        "answer": original,
                        "shown": shown,
                        "correct": original == question["correct"],
                    }
                else:
                    responses[str(item["qid"])] = {"answer": None, "shown": None, "correct": None}
            entry.update(
                {
                    "started_at": attempt["started_at"],
                    "deadline": attempt["deadline"],
                    "answered": len(answers),
                    "total": len(items),
                    "resume_allowed": bool(attempt["resume_allowed"]),
                    "responses": responses,
                }
            )
            if attempt["submitted_at"] is None:
                entry["state"] = "in_progress"
                entry["remaining_seconds"] = max(0, int(attempt["deadline"] - time.time()))
            else:
                correct = attempt["correct_count"] or 0
                total = attempt["total"] or len(items)
                entry.update(
                    {
                        "state": "submitted",
                        "correct": correct,
                        "total": total,
                        "score10": exam_score10(correct, total),
                        "submitted_at": attempt["submitted_at"],
                        "reason": attempt["submit_reason"],
                        "seconds_used": max(0, int(attempt["submitted_at"] - attempt["started_at"])),
                    }
                )
                score_total += entry["score10"]
        summary[entry["state"]] += 1
        rows.append(entry)
    summary["average_score10"] = round(score_total / summary["submitted"], 2) if summary["submitted"] else None
    return rows, summary, questions


def build_exam_analysis(questions, rows):
    """Phân tích từng câu trên các bài ĐÃ NỘP (bài đang làm dở chưa tính để không bị lệch vì câu chưa kịp trả lời).

    - tỷ lệ sai = (chọn sai + bỏ trống) / số sinh viên được ra câu đó — khớp cách chấm điểm (bỏ trống tính là sai).
    - top_wrong: đáp án SAI được chọn nhiều nhất (chữ A-D theo ngân hàng câu hỏi); nếu đồng số lượt thì liệt kê tất cả.
    Kết quả xếp theo tỷ lệ sai giảm dần.
    """
    def empty_bucket():
        bucket = {letter: 0 for letter in EXAM_LETTERS}
        bucket.update({"blank": 0, "taken": 0})
        return bucket

    stats = {question["id"]: empty_bucket() for question in questions}
    for entry in rows:
        if entry["state"] != "submitted":
            continue
        for question_id, response in entry["responses"].items():
            bucket = stats.get(int(question_id))
            if bucket is None:
                continue
            bucket["taken"] += 1
            if response["answer"] in EXAM_LETTERS:
                bucket[response["answer"]] += 1
            else:
                bucket["blank"] += 1

    analysis = []
    for question in questions:
        bucket = stats[question["id"]]
        taken = bucket["taken"]
        if taken == 0:
            continue
        correct_letter = question["correct"]
        right = bucket[correct_letter] if correct_letter in EXAM_LETTERS else 0
        wrong_choices = {letter: bucket[letter] for letter in EXAM_LETTERS if letter != correct_letter and bucket[letter] > 0}
        wrong_answered = sum(wrong_choices.values())
        wrong_total = wrong_answered + bucket["blank"]
        top_count = max(wrong_choices.values()) if wrong_choices else 0
        analysis.append(
            {
                "id": question["id"],
                "position": question["position"],
                "text": question["text"],
                "correct": correct_letter,
                "taken": taken,
                "right": right,
                "wrong_answered": wrong_answered,
                "blank": bucket["blank"],
                "wrong_total": wrong_total,
                "error_rate": round(wrong_total / taken, 4),
                "distribution": {letter: bucket[letter] for letter in EXAM_LETTERS},
                "top_wrong": [
                    {"letter": letter, "count": count, "text": question["options"].get(letter, "")}
                    for letter, count in wrong_choices.items()
                    if count == top_count
                ],
            }
        )
    analysis.sort(key=lambda item: (-item["error_rate"], -item["wrong_total"], item["position"]))
    for rank, item in enumerate(analysis, start=1):
        item["rank"] = rank
    return analysis


@app.route("/api/rooms/exams", methods=["GET", "POST"])
def room_exams():
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên trong phòng mới quản lý được bài kiểm tra."}), 403
    room_code = session["room_code"]
    conn = get_db_connection()
    bank_size = conn.execute(
        "SELECT COUNT(*) FROM questions WHERE room_code = ? AND presentation_id IS NULL",
        (room_code,),
    ).fetchone()[0]

    if request.method == "POST":
        data = get_json_data()
        title = str(data.get("title", "")).strip()[:120] or "Bài kiểm tra"
        try:
            duration = int(data.get("duration_minutes"))
            count = int(data.get("question_count"))
        except (TypeError, ValueError):
            conn.close()
            return jsonify({"status": "error", "message": "Thời gian và số câu phải là số nguyên."}), 400
        if bank_size == 0:
            conn.close()
            return jsonify({"status": "error", "message": "Ngân hàng câu hỏi đang trống. Hãy thêm câu hỏi trước."}), 400
        if not 1 <= duration <= EXAM_MAX_MINUTES:
            conn.close()
            return jsonify({"status": "error", "message": f"Thời gian làm bài từ 1 đến {EXAM_MAX_MINUTES} phút."}), 400
        requested_question_ids = data.get("question_ids")
        if requested_question_ids is None:
            question_pool_ids = [
                row["id"] for row in conn.execute(
                    "SELECT id FROM questions WHERE room_code = ? AND presentation_id IS NULL",
                    (room_code,),
                )
            ]
        else:
            if not isinstance(requested_question_ids, list) or any(
                isinstance(question_id, bool) or not isinstance(question_id, int)
                for question_id in requested_question_ids
            ):
                conn.close()
                return jsonify({"status": "error", "message": "Danh sách câu hỏi được chọn không hợp lệ."}), 400
            question_pool_ids = list(dict.fromkeys(requested_question_ids))
            if not question_pool_ids:
                conn.close()
                return jsonify({"status": "error", "message": "Hãy chọn ít nhất một câu hỏi cho bài kiểm tra."}), 400
            placeholders = ",".join("?" for _ in question_pool_ids)
            valid_question_ids = {
                row["id"] for row in conn.execute(
                    "SELECT id FROM questions WHERE room_code = ? AND presentation_id IS NULL "
                    f"AND id IN ({placeholders})",
                    [room_code, *question_pool_ids],
                )
            }
            if len(valid_question_ids) != len(question_pool_ids):
                conn.close()
                return jsonify({
                    "status": "error",
                    "message": "Một hoặc nhiều câu hỏi đã chọn không còn trong ngân hàng của phòng này.",
                }), 400
        if not 1 <= count <= len(question_pool_ids):
            conn.close()
            return jsonify({
                "status": "error",
                "message": f"Số câu cần bốc phải từ 1 đến {len(question_pool_ids)} câu đã chọn.",
            }), 400
        try:
            scheduled_at = normalize_exam_date(data.get("exam_date"))
        except ValueError as error:
            conn.close()
            return jsonify({"status": "error", "message": str(error)}), 400
        access_code = generate_exam_code(conn)
        cursor = conn.execute(
            "INSERT INTO exams (room_code, title, duration_minutes, question_count, access_code, status, "
            "show_answers, created_at, scheduled_at, question_pool_json) "
            "VALUES (?, ?, ?, ?, ?, 'draft', ?, ?, ?, ?)",
            (
                room_code,
                title,
                duration,
                count,
                access_code,
                int(bool(data.get("show_answers"))),
                time.time(),
                scheduled_at,
                json.dumps(question_pool_ids),
            ),
        )
        conn.commit()
        exam = conn.execute("SELECT * FROM exams WHERE id = ?", (cursor.lastrowid,)).fetchone()
        conn.close()
        return jsonify({"status": "success", "exam": serialize_exam(exam)}), 201

    exams = conn.execute("SELECT * FROM exams WHERE room_code = ? ORDER BY id DESC", (room_code,)).fetchall()
    class_size = conn.execute("SELECT COUNT(*) FROM room_students WHERE room_code = ?", (room_code,)).fetchone()[0]
    result = []
    for exam in exams:
        finalize_due_attempts(conn, exam)
        counts = conn.execute(
            "SELECT COUNT(*) AS started, COALESCE(SUM(submitted_at IS NOT NULL), 0) AS submitted "
            "FROM exam_attempts WHERE exam_id = ?",
            (exam["id"],),
        ).fetchone()
        item = serialize_exam(exam)
        item.update({"started": counts["started"], "submitted": counts["submitted"]})
        result.append(item)
    bank_questions = conn.execute(
        "SELECT id, position, text FROM questions WHERE room_code = ? "
        "AND presentation_id IS NULL ORDER BY position",
        (room_code,),
    ).fetchall()
    conn.close()
    return jsonify({
        "status": "success",
        "exams": result,
        "class_size": class_size,
        "bank_size": bank_size,
        "question_bank": [dict(row) for row in bank_questions],
    })


@app.route("/api/rooms/exams/<int:exam_id>/status", methods=["POST"])
def set_exam_status(exam_id):
    """draft -> open (cho sinh viên vào làm) -> closed (kết thúc, tự nộp mọi bài đang làm)."""
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên trong phòng mới đổi được trạng thái bài kiểm tra."}), 403
    target = str(get_json_data().get("status", "")).strip()
    conn = get_db_connection()
    exam = get_teacher_exam(conn, exam_id)
    if exam is None:
        conn.close()
        return exam_not_found_response()
    allowed = {"draft": "open", "open": "closed"}
    if target not in {"open", "closed"} or allowed.get(exam["status"]) != target:
        conn.close()
        return jsonify({"status": "error", "message": "Không thể chuyển bài kiểm tra sang trạng thái này."}), 409
    if target == "open":
        conn.execute("UPDATE exams SET status = 'open', opened_at = ? WHERE id = ?", (time.time(), exam_id))
    else:
        conn.execute("UPDATE exams SET status = 'closed', closed_at = ? WHERE id = ?", (time.time(), exam_id))
    conn.commit()
    if target == "closed":
        finalize_due_attempts(conn, conn.execute("SELECT * FROM exams WHERE id = ?", (exam_id,)).fetchone())
    conn.close()
    return jsonify({"status": "success"})


@app.route("/api/rooms/exams/<int:exam_id>/schedule", methods=["POST"])
def set_exam_schedule(exam_id):
    """Đặt/đổi/xóa ngày kiểm tra của một bài (kể cả bài đã tạo từ trước)."""
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên trong phòng mới đổi được ngày kiểm tra."}), 403
    try:
        scheduled_at = normalize_exam_date(get_json_data().get("exam_date"))
    except ValueError as error:
        return jsonify({"status": "error", "message": str(error)}), 400
    conn = get_db_connection()
    if get_teacher_exam(conn, exam_id) is None:
        conn.close()
        return exam_not_found_response()
    conn.execute("UPDATE exams SET scheduled_at = ? WHERE id = ?", (scheduled_at, exam_id))
    conn.commit()
    conn.close()
    return jsonify({"status": "success", "scheduled_at": scheduled_at})


@app.route("/api/rooms/exams/<int:exam_id>", methods=["DELETE"])
def delete_exam(exam_id):
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên trong phòng mới xóa được bài kiểm tra."}), 403
    conn = get_db_connection()
    if get_teacher_exam(conn, exam_id) is None:
        conn.close()
        return exam_not_found_response()
    conn.execute(
        "DELETE FROM exam_answers WHERE attempt_id IN (SELECT id FROM exam_attempts WHERE exam_id = ?)", (exam_id,)
    )
    conn.execute("DELETE FROM exam_attempts WHERE exam_id = ?", (exam_id,))
    conn.execute("DELETE FROM exams WHERE id = ?", (exam_id,))
    conn.commit()
    conn.close()
    return jsonify({"status": "success"})


@app.route("/api/rooms/exams/<int:exam_id>/results", methods=["GET"])
def get_exam_results(exam_id):
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên trong phòng mới xem được kết quả."}), 403
    conn = get_db_connection()
    exam = get_teacher_exam(conn, exam_id)
    if exam is None:
        conn.close()
        return exam_not_found_response()
    rows, summary, questions = build_exam_results(conn, exam)
    exam = get_teacher_exam(conn, exam_id)
    conn.close()
    return jsonify(
        {
            "status": "success",
            "exam": serialize_exam(exam),
            "summary": summary,
            "students": rows,
            "questions": questions,
            "analysis": build_exam_analysis(questions, rows),
        }
    )


@app.route("/api/rooms/exams/<int:exam_id>/students/reset", methods=["POST"])
def reset_exam_student(exam_id):
    """resume: cho sinh viên đang làm dở vào lại (đồng hồ vẫn chạy); delete: xóa bài để làm lại từ đầu."""
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên trong phòng mới thao tác được."}), 403
    data = get_json_data()
    student_id = clean_student_id(data.get("student_id"))
    mode = str(data.get("mode", "")).strip()
    if not student_id or mode not in {"resume", "delete"}:
        return jsonify({"status": "error", "message": "Thiếu MSSV hoặc thao tác không hợp lệ."}), 400
    conn = get_db_connection()
    exam = get_teacher_exam(conn, exam_id)
    if exam is None:
        conn.close()
        return exam_not_found_response()
    attempt = conn.execute(
        "SELECT * FROM exam_attempts WHERE exam_id = ? AND student_id = ?", (exam_id, student_id)
    ).fetchone()
    if attempt is None:
        conn.close()
        return jsonify({"status": "error", "message": "Sinh viên này chưa có bài làm."}), 404
    attempt = auto_finalize(conn, attempt, exam)
    if mode == "resume":
        if attempt["submitted_at"] is not None:
            conn.close()
            return jsonify(
                {"status": "error", "message": "Bài đã được nộp. Hãy chọn “Xóa bài” nếu muốn cho làm lại."}
            ), 409
        conn.execute("UPDATE exam_attempts SET resume_allowed = 1 WHERE id = ?", (attempt["id"],))
        message = "Đã cho phép vào lại. Sinh viên nhập lại mã bài và MSSV để tiếp tục (đồng hồ vẫn chạy)."
    else:
        conn.execute("DELETE FROM exam_answers WHERE attempt_id = ?", (attempt["id"],))
        conn.execute("DELETE FROM exam_attempts WHERE id = ?", (attempt["id"],))
        message = "Đã xóa bài làm. Sinh viên có thể vào làm lại với bộ câu hỏi mới."
    conn.commit()
    conn.close()
    return jsonify({"status": "success", "message": message})


def format_local_time(timestamp):
    if not timestamp:
        return ""
    return datetime.fromtimestamp(timestamp).strftime("%d/%m/%Y %H:%M:%S")


@app.route("/api/rooms/exams/<int:exam_id>/results.xlsx", methods=["GET"])
def export_exam_results(exam_id):
    """Xuất kết quả dạng MSSV x câu hỏi: đáp án đúng nền xanh, sai nền đỏ, chưa trả lời nền xám."""
    if not teacher_room_required():
        return jsonify({"status": "error", "message": "Chỉ giảng viên trong phòng mới xuất được kết quả."}), 403
    conn = get_db_connection()
    exam = get_teacher_exam(conn, exam_id)
    if exam is None:
        conn.close()
        return exam_not_found_response()
    rows, summary, questions = build_exam_results(conn, exam)
    exam = get_teacher_exam(conn, exam_id)
    conn.close()

    from openpyxl.utils import get_column_letter

    state_labels = {"not_started": "Chưa làm", "in_progress": "Đang làm", "submitted": "Đã nộp"}
    reason_labels = {"manual": "Tự nộp", "timeout": "Hết giờ", "closed": "Giảng viên kết thúc"}

    def spreadsheet_text(value):
        value = "" if value is None else str(value)
        return "'" + value if value.startswith(("=", "+", "-", "@")) else value

    def minutes_seconds(total_seconds):
        return f"{total_seconds // 60:02d}:{total_seconds % 60:02d}"

    header_fill = PatternFill("solid", fgColor="1E3A8A")
    green_fill, green_font = PatternFill("solid", fgColor="DCFCE7"), Font(color="166534", bold=True)
    red_fill, red_font = PatternFill("solid", fgColor="FEE2E2"), Font(color="991B1B", bold=True)
    gray_fill, gray_font = PatternFill("solid", fgColor="F1F5F9"), Font(color="64748B", italic=True)

    def style_header(row_cells):
        for cell in row_cells:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Kết quả"
    sheet.append([spreadsheet_text(exam["title"])])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append(
        [
            f"Số câu/đề: {exam['question_count']}",
            f"Thời gian làm bài: {exam['duration_minutes']} phút",
            f"Mã bài: {exam['access_code']}",
        ]
    )
    sheet.append(
        [
            f"Ngày kiểm tra: {format_exam_date(exam['scheduled_at']) or '—'}",
            f"Mở bài lúc: {format_local_time(exam['opened_at']) or '—'}",
            f"Kết thúc lúc: {format_local_time(exam['closed_at']) or '—'}",
        ]
    )
    sheet.append([])
    question_count = len(questions)
    headers = (
        ["STT", "MSSV", "Họ và tên", "Trạng thái"]
        + [f"Câu {question['position']}" for question in questions]
        + ["Số câu đúng", "Tổng số câu", "Điểm (thang 10)", "Bắt đầu", "Nộp lúc", "Thời gian làm (mm:ss)", "Hình thức nộp"]
    )
    sheet.append(headers)
    header_row = sheet.max_row
    style_header(sheet[header_row])
    sheet.row_dimensions[header_row].height = 30

    first_question_col = 5
    stats_col = first_question_col + question_count  # cột "Số câu đúng"
    for entry in rows:
        submitted = entry["state"] == "submitted"
        output_row = [
            entry["roll_number"] or "",
            spreadsheet_text(entry["student_id"]),
            spreadsheet_text(entry["name"]),
            state_labels[entry["state"]],
        ]
        cell_styles = []
        for question in questions:
            response = entry["responses"].get(str(question["id"]))
            if response is None:
                output_row.append("")  # câu này không nằm trong đề của sinh viên
                cell_styles.append(None)
            elif response["answer"] is None:
                output_row.append("—")
                cell_styles.append("blank")
            else:
                output_row.append(response["answer"])
                cell_styles.append("correct" if response["correct"] else "wrong")
        output_row += [
            entry.get("correct", "") if submitted else "",
            entry.get("total", "") if entry["state"] != "not_started" else "",
            entry.get("score10", "") if submitted else "",
            format_local_time(entry.get("started_at")),
            format_local_time(entry.get("submitted_at")) if submitted else "",
            minutes_seconds(entry["seconds_used"]) if submitted else "",
            reason_labels.get(entry.get("reason"), "") if submitted else "",
        ]
        sheet.append(output_row)
        row_index = sheet.max_row
        for offset, style in enumerate(cell_styles):
            cell = sheet.cell(row=row_index, column=first_question_col + offset)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if style == "correct":
                cell.fill, cell.font = green_fill, green_font
            elif style == "wrong":
                cell.fill, cell.font = red_fill, red_font
            elif style == "blank":
                cell.fill, cell.font = gray_fill, gray_font
        if submitted:
            score_cell = sheet.cell(row=row_index, column=stats_col + 2)
            passed = entry["score10"] >= 5
            score_cell.fill, score_cell.font = (green_fill, green_font) if passed else (red_fill, red_font)

    sheet.append([])
    sheet.append(
        [
            "", "", "Sĩ số lớp", summary["class_size"], "Đã nộp", summary["submitted"],
            "Điểm TB", summary["average_score10"] if summary["average_score10"] is not None else "",
        ]
    )
    for column, width in zip("ABCD", (7, 16, 28, 12)):
        sheet.column_dimensions[column].width = width
    for index in range(question_count):
        sheet.column_dimensions[get_column_letter(first_question_col + index)].width = 9
    for offset, width in enumerate((13, 13, 16, 20, 20, 22, 20)):
        sheet.column_dimensions[get_column_letter(stats_col + offset)].width = width
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=first_question_col)

    # Trang 2: nội dung câu hỏi, đáp án đúng và tỷ lệ làm đúng để đối chiếu với các cột "Câu N".
    detail = workbook.create_sheet("Đề và đáp án")
    detail.append(
        ["Câu", "Nội dung", "A", "B", "C", "D", "Đáp án đúng", "Số SV được ra đề", "Số SV trả lời đúng", "Tỷ lệ đúng"]
    )
    style_header(detail[1])
    for question in questions:
        taken = right = 0
        for entry in rows:
            response = entry["responses"].get(str(question["id"]))
            if response is not None:
                taken += 1
                right += 1 if response["correct"] else 0
        detail.append(
            [
                question["position"],
                spreadsheet_text(question["text"]),
                *[spreadsheet_text(question["options"].get(letter, "")) for letter in EXAM_LETTERS],
                question["correct"],
                taken,
                right,
                right / taken if taken else "",
            ]
        )
        row_index = detail.max_row
        correct_cell = detail.cell(row=row_index, column=7)
        correct_cell.fill, correct_cell.font = green_fill, green_font
        correct_cell.alignment = Alignment(horizontal="center", vertical="top")
        percent_cell = detail.cell(row=row_index, column=10)
        percent_cell.number_format = "0.0%"
        if taken:
            passed = right / taken >= 0.5
            percent_cell.fill, percent_cell.font = (green_fill, green_font) if passed else (red_fill, red_font)
    for column, width in zip("ABCDEFGHIJ", (7, 48, 26, 26, 26, 26, 13, 17, 18, 13)):
        detail.column_dimensions[column].width = width
    detail.freeze_panes = "C2"
    for row in detail.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top", wrap_text=True, horizontal=cell.alignment.horizontal
            )

    # Trang 3: phân tích câu hỏi — câu nào sai nhiều nhất, đáp án sai nào được chọn nhiều nhất.
    analysis = build_exam_analysis(questions, rows)
    amber_fill, amber_font = PatternFill("solid", fgColor="FEF3C7"), Font(color="92400E", bold=True)
    report = workbook.create_sheet("Phân tích câu hỏi")
    report.append(
        [
            "Hạng", "Câu", "Nội dung", "Đáp án đúng", "Số SV làm", "Đúng", "Chọn sai", "Bỏ trống",
            "Tỷ lệ sai", "Chọn A", "Chọn B", "Chọn C", "Chọn D",
            "Đáp án sai được chọn nhiều nhất", "Số lượt chọn", "Nội dung đáp án sai đó",
        ]
    )
    style_header(report[1])
    report.row_dimensions[1].height = 32
    for item in analysis:
        top_letters = ", ".join(entry["letter"] for entry in item["top_wrong"])
        top_count = item["top_wrong"][0]["count"] if item["top_wrong"] else ""
        top_text = " | ".join(f"{entry['letter']}: {entry['text']}" for entry in item["top_wrong"])
        report.append(
            [
                item["rank"], item["position"], spreadsheet_text(item["text"]), item["correct"], item["taken"],
                item["right"], item["wrong_answered"], item["blank"], item["error_rate"],
                *[item["distribution"][letter] for letter in EXAM_LETTERS],
                top_letters or "—", top_count, spreadsheet_text(top_text),
            ]
        )
        row_index = report.max_row
        rate_cell = report.cell(row=row_index, column=9)
        rate_cell.number_format = "0.0%"
        if item["error_rate"] >= 0.5:
            rate_cell.fill, rate_cell.font = red_fill, red_font
        elif item["error_rate"] >= 0.3:
            rate_cell.fill, rate_cell.font = amber_fill, amber_font
        else:
            rate_cell.fill, rate_cell.font = green_fill, green_font
        correct_letter_cell = report.cell(row=row_index, column=4)
        correct_letter_cell.fill, correct_letter_cell.font = green_fill, green_font
        correct_letter_cell.alignment = Alignment(horizontal="center", vertical="top")
        for letter_index, letter in enumerate(EXAM_LETTERS):
            count_cell = report.cell(row=row_index, column=10 + letter_index)
            count_cell.alignment = Alignment(horizontal="center", vertical="top")
            if letter == item["correct"]:
                count_cell.fill, count_cell.font = green_fill, green_font
            elif any(entry["letter"] == letter for entry in item["top_wrong"]):
                count_cell.fill, count_cell.font = red_fill, red_font
        wrong_cell = report.cell(row=row_index, column=14)
        wrong_cell.alignment = Alignment(horizontal="center", vertical="top")
        if item["top_wrong"]:
            wrong_cell.fill, wrong_cell.font = red_fill, red_font
    if not analysis:
        report.append(["Chưa có bài nộp nào để phân tích."])
    else:
        report.append([])
        report.append(
            ["Ghi chú: chỉ tính các bài đã nộp; câu bỏ trống tính là sai; chữ A–D theo ngân hàng câu hỏi; "
             "xếp theo tỷ lệ sai giảm dần."]
        )
    for column, width in zip("ABCDEFGHIJKLMNOP", (7, 7, 48, 12, 11, 9, 10, 10, 11, 9, 9, 9, 9, 20, 13, 40)):
        report.column_dimensions[column].width = width
    report.freeze_panes = "D2"
    for row in report.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True, horizontal=cell.alignment.horizontal)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name=f"Kiem_Tra_{session['room_code']}_{exam['access_code']}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ----- API sinh viên (không cần đăng nhập; xác thực bằng mã bài + MSSV, sau đó bằng token) -----
def token_matches(attempt, token):
    """True nếu token do trình duyệt gửi lên đúng là token của lượt làm bài này."""
    token = str(token or "").strip()
    return bool(token) and secrets.compare_digest(attempt["token"], token)


def lookup_exam_and_student(conn, data):
    """Kiểm tra mã bài và MSSV. Trả về (exam, student, error_response)."""
    access_code = str(data.get("access_code", "")).strip().upper()
    student_id = clean_student_id(data.get("student_id"))
    if not access_code or not student_id:
        return None, None, (jsonify({"status": "error", "message": "Vui lòng nhập mã bài kiểm tra và MSSV."}), 400)
    exam = conn.execute("SELECT * FROM exams WHERE access_code = ?", (access_code,)).fetchone()
    if exam is None:
        return None, None, (jsonify({"status": "error", "message": "Mã bài kiểm tra không đúng."}), 404)
    student = conn.execute(
        "SELECT student_id, name FROM room_students WHERE room_code = ? AND student_id = ?",
        (exam["room_code"], student_id),
    ).fetchone()
    if student is None:
        return None, None, (
            jsonify({"status": "error", "message": "MSSV không có trong danh sách lớp của bài kiểm tra này."}),
            404,
        )
    return exam, student, None


@app.route("/api/exam/lookup", methods=["POST"])
def exam_lookup():
    """Bước xác nhận trước khi tính giờ: hiện tên sinh viên và thông tin bài."""
    conn = get_db_connection()
    data = get_json_data()
    exam, student, error = lookup_exam_and_student(conn, data)
    if error:
        conn.close()
        return error
    attempt = conn.execute(
        "SELECT * FROM exam_attempts WHERE exam_id = ? AND student_id = ?", (exam["id"], student["student_id"])
    ).fetchone()
    if attempt is not None:
        attempt = auto_finalize(conn, attempt, exam)
    available_question_count = len(get_exam_question_ids(conn, exam))
    conn.close()

    if attempt is None:
        attempt_state = "new"
    else:
        attempt_state = "submitted" if attempt["submitted_at"] is not None else "in_progress"
    message = ""
    if attempt_state == "new":
        if exam["status"] == "draft":
            message = "Bài kiểm tra chưa được mở. Vui lòng chờ giảng viên."
        elif exam["status"] == "closed":
            message = "Bài kiểm tra đã kết thúc."
    return jsonify(
        {
            "status": "success",
            "exam": {
                "title": exam["title"],
                "duration_minutes": exam["duration_minutes"],
                "question_count": min(exam["question_count"], available_question_count),
                "status": exam["status"],
                "scheduled_at": exam["scheduled_at"],
            },
            "student": {"student_id": student["student_id"], "name": student["name"]},
            "attempt_state": attempt_state,
            "can_resume": bool(
                attempt is not None
                and attempt["submitted_at"] is None
                and (attempt["resume_allowed"] or token_matches(attempt, data.get("resume_token")))
            ),
            "can_start": attempt_state == "new" and exam["status"] == "open",
            "message": message,
        }
    )


@app.route("/api/exam/start", methods=["POST"])
def exam_start():
    """Bắt đầu tính giờ: bốc ngẫu nhiên câu hỏi, xáo thứ tự câu và thứ tự đáp án riêng cho sinh viên này."""
    conn = get_db_connection()
    data = get_json_data()
    exam, student, error = lookup_exam_and_student(conn, data)
    if error:
        conn.close()
        return error

    attempt = conn.execute(
        "SELECT * FROM exam_attempts WHERE exam_id = ? AND student_id = ?", (exam["id"], student["student_id"])
    ).fetchone()
    if attempt is not None:
        attempt = auto_finalize(conn, attempt, exam)
        if attempt["submitted_at"] is not None:
            conn.close()
            return jsonify({"status": "error", "message": "Bạn đã nộp bài kiểm tra này rồi."}), 409
        own_session = token_matches(attempt, data.get("resume_token"))
        if not attempt["resume_allowed"] and not own_session:
            conn.close()
            return jsonify(
                {
                    "status": "error",
                    "message": "Bài làm của bạn đang được mở ở một thiết bị/trình duyệt khác. "
                    "Nếu bị mất kết nối, hãy báo giảng viên chọn “Cho vào lại”.",
                }
            ), 409
        if own_session:
            # Chính trình duyệt này đã bắt đầu bài (ví dụ lỡ tải lại trang): vào lại, giữ nguyên phiên và đồng hồ.
            payload = build_attempt_payload(conn, attempt, exam, include_token=True)
            conn.close()
            return jsonify(payload)
        # Vào lại có sự cho phép của giảng viên: đổi token để vô hiệu hóa phiên cũ.
        new_token = secrets.token_urlsafe(24)
        conn.execute(
            "UPDATE exam_attempts SET token = ?, resume_allowed = 0 WHERE id = ?", (new_token, attempt["id"])
        )
        conn.commit()
        attempt = refresh_attempt(conn, attempt["id"])
        payload = build_attempt_payload(conn, attempt, exam, include_token=True)
        conn.close()
        return jsonify(payload)

    if exam["status"] != "open":
        conn.close()
        message = "Bài kiểm tra chưa được mở. Vui lòng chờ giảng viên." if exam["status"] == "draft" else "Bài kiểm tra đã kết thúc."
        return jsonify({"status": "error", "message": message}), 409

    question_ids = get_exam_question_ids(conn, exam)
    count = min(exam["question_count"], len(question_ids))
    if count < 1:
        conn.close()
        return jsonify({"status": "error", "message": "Nhóm câu hỏi đã chọn không còn câu hợp lệ trong ngân hàng."}), 409

    # sample() vừa chọn ngẫu nhiên vừa cho thứ tự ngẫu nhiên; mỗi câu lại xáo thứ tự A-D riêng.
    chosen_ids = _exam_rng.sample(question_ids, count)
    items = []
    for question_id in chosen_ids:
        order = list(EXAM_LETTERS)
        _exam_rng.shuffle(order)
        items.append({"qid": question_id, "order": order})

    started_at = time.time()
    try:
        cursor = conn.execute(
            "INSERT INTO exam_attempts (exam_id, student_id, token, started_at, deadline, items_json) "
            "VALUES (?, ?, ?, ?, ?, ?)",
            (
                exam["id"],
                student["student_id"],
                secrets.token_urlsafe(24),
                started_at,
                started_at + exam["duration_minutes"] * 60,
                json.dumps(items),
            ),
        )
        conn.commit()
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({"status": "error", "message": "Bài làm đã được tạo trước đó. Hãy thử lại."}), 409
    attempt = refresh_attempt(conn, cursor.lastrowid)
    payload = build_attempt_payload(conn, attempt, exam, include_token=True)
    conn.close()
    return jsonify(payload), 201


@app.route("/api/exam/attempt", methods=["GET"])
def exam_get_attempt():
    """Khôi phục bài đang làm (sau khi tải lại trang) và đồng bộ đồng hồ với máy chủ."""
    conn = get_db_connection()
    attempt, exam = load_attempt_from_token(conn)
    if attempt is None:
        conn.close()
        return invalid_attempt_response()
    attempt = auto_finalize(conn, attempt, exam)
    payload = build_attempt_payload(conn, attempt, exam)
    conn.close()
    return jsonify(payload)


@app.route("/api/exam/answer", methods=["POST"])
def exam_save_answer():
    """Lưu đáp án ngay khi sinh viên chọn (letter là chữ cái đang hiển thị trên màn hình sinh viên)."""
    conn = get_db_connection()
    attempt, exam = load_attempt_from_token(conn)
    if attempt is None:
        conn.close()
        return invalid_attempt_response()
    attempt = auto_finalize(conn, attempt, exam)
    if attempt["submitted_at"] is not None:
        conn.close()
        return jsonify(
            {"status": "error", "code": "submitted", "message": "Bài đã được nộp hoặc đã hết giờ; đáp án không được ghi nhận."}
        ), 409

    data = get_json_data()
    try:
        question_id = int(data.get("qid"))
    except (TypeError, ValueError):
        conn.close()
        return jsonify({"status": "error", "message": "Câu hỏi không hợp lệ."}), 400
    letter = str(data.get("letter") or "").strip().upper()
    if question_id not in {item["qid"] for item in json.loads(attempt["items_json"])}:
        conn.close()
        return jsonify({"status": "error", "message": "Câu hỏi không thuộc bài làm của bạn."}), 400
    if letter and letter not in EXAM_LETTERS:
        conn.close()
        return jsonify({"status": "error", "message": "Đáp án phải là A, B, C hoặc D."}), 400

    if letter:
        conn.execute(
            "INSERT INTO exam_answers (attempt_id, question_id, choice, answered_at) VALUES (?, ?, ?, ?) "
            "ON CONFLICT(attempt_id, question_id) DO UPDATE SET choice = excluded.choice, answered_at = excluded.answered_at",
            (attempt["id"], question_id, letter, time.time()),
        )
    else:
        conn.execute(
            "DELETE FROM exam_answers WHERE attempt_id = ? AND question_id = ?", (attempt["id"], question_id)
        )
    conn.commit()
    remaining = max(0, int(attempt["deadline"] - time.time()))
    conn.close()
    return jsonify({"status": "success", "remaining_seconds": remaining})


@app.route("/api/exam/submit", methods=["POST"])
def exam_submit():
    conn = get_db_connection()
    attempt, exam = load_attempt_from_token(conn)
    if attempt is None:
        conn.close()
        return invalid_attempt_response()
    attempt = auto_finalize(conn, attempt, exam)
    if attempt["submitted_at"] is None:
        finalize_attempt(conn, attempt, "manual")
        conn.commit()
        attempt = refresh_attempt(conn, attempt["id"])
    payload = build_attempt_payload(conn, attempt, exam)
    conn.close()
    return jsonify(payload)


# -----------------------------------------------------------------------------
# Khởi chạy Ứng dụng
# -----------------------------------------------------------------------------
if __name__ == "__main__":
    # Development server chỉ dùng để kiểm thử. Production chạy WSGI server
    # (Gunicorn) phía sau reverse proxy HTTPS như Nginx/Cloudflare.
    serial_thread = threading.Thread(target=read_serial_data, daemon=True)
    serial_thread.start()
    mqtt_client = start_mqtt_client()
    app.run(
        host=os.environ.get("FLASK_HOST", "127.0.0.1" if IS_PRODUCTION else "0.0.0.0"),
        port=int(os.environ.get("FLASK_PORT", "5000")),
        debug=False,
        use_reloader=False,
    )
